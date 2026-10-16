from pathlib import Path
from zipfile import ZipFile
from lxml import etree as E
from docx import Document
from collections import Counter
import json,re,hashlib,sys
p=Path(sys.argv[1]);f=p/'گزارش_پروژه_کارشناسی_ایمان_کربلائی.docx'
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
W='{'+ns['w']+'}'
with ZipFile(f) as z:
 bad=z.testzip();xml=E.fromstring(z.read('word/document.xml'));styles=E.fromstring(z.read('word/styles.xml'));settings=E.fromstring(z.read('word/settings.xml'))
 duplicates=[]
 for name,tree in [('document',xml),('styles',styles)]:
  for e in tree.xpath('//w:rPr|//w:pPr|//w:sectPr',namespaces=ns):
   for tag,c in Counter((x.tag,x.get(W+'type')) if x.tag in [W+'headerReference',W+'footerReference'] else (x.tag,None) for x in e).items():
    if c>1:duplicates.append((name,tag,c))
 d=Document(f);paras=d.paragraphs
 h1=[x.text for x in paras if x.style.name=='Heading 1'];h2=[x.text for x in paras if x.style.name=='Heading 2']
 wrong=[]
 for i,x in enumerate(paras):
  if re.search('[\u0600-\u06ff]',x.text) and x.style.name not in ['Code']:
   direct=x._p.find(W+'pPr/'+W+'bidi')
   inherited=False
   style=x.style
   while style is not None:
    pr=style._element.pPr
    node=pr.find(W+'bidi') if pr is not None else None
    if node is not None:
     inherited=node.get(W+'val')!='0'
     break
    style=style.base_style
   if (direct is not None and direct.get(W+'val')=='0') or (direct is None and not inherited):wrong.append(i)
 run_count=len(xml.xpath('//w:r',namespaces=ns));fa_runs=0;ltr=0
 for x in xml.xpath('//w:r',namespaces=ns):
  v=''.join(x.xpath('.//w:t/text()',namespaces=ns))
  if re.search('[\u0600-\u06ff]',v):fa_runs+=1
  if x.xpath('./w:rPr/w:rtl[@w:val="0"]',namespaces=ns):ltr+=1
 b=xml.xpath('//w:bookmarkStart[@w:name="BismillahReserved"]',namespaces=ns)[0].getparent()
 text=''.join(b.xpath('.//w:t/text()',namespaces=ns))
 sections=[]
 for i,s in enumerate(d.sections):
  sections.append({'index':i,'start_type':str(s.start_type),'width_cm':round(s.page_width.cm,3),'height_cm':round(s.page_height.cm,3),'top_cm':round(s.top_margin.cm,3),'bottom_cm':round(s.bottom_margin.cm,3),'left_odd_cm':round(s.left_margin.cm,3),'right_odd_cm':round(s.right_margin.cm,3),'different_first':s.different_first_page_header_footer})
 lines=(p/'report-text-rebuilt.md').read_text('utf8').splitlines();dupbody=Counter(x for x in lines if len(x)>180 and not x.startswith('مبنای'))
 result={'file':str(f),'bytes':f.stat().st_size,'sha256':hashlib.sha256(f.read_bytes()).hexdigest(),'zip_integrity_error':bad,'duplicate_property_nodes':duplicates,'main_chapters':len([x for x in h1 if x.startswith('فصل')]),'sections':sections,'tables':len(d.tables),'figures':len(d.inline_shapes),'heading2_count':len(h2),'main_text_words':len((p/'report-text-rebuilt.md').read_text('utf8').split()),'total_paragraphs':len(paras),'rtl_persian_paragraph_failures':wrong,'persian_runs':fa_runs,'ltr_runs':ltr,'total_runs':run_count,'bismillah_paragraph_text':text,'update_fields':bool(settings.xpath('//w:updateFields[@w:val="1"]',namespaces=ns)),'mirror_margins':bool(settings.xpath('//w:mirrorMargins',namespaces=ns)),'toc_field_codes':xml.xpath('//w:instrText/text()',namespaces=ns),'duplicate_long_body_paragraphs':[x[:90] for x,c in dupbody.items() if c>1],'actual_rendered_pages':None,'actual_main_body_pages':None,'rendered_visual_inspection':False}
(p/'structural-audit.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8');print(json.dumps({k:v for k,v in result.items() if k not in ['sections','toc_field_codes']},ensure_ascii=True))
