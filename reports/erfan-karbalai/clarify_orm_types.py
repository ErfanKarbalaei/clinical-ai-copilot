"""Clarify generic ORM type display without touching application sources."""
from pathlib import Path
from docx import Document
out=Path(__file__).resolve().parent
path=out/'گزارش_پروژه_کارشناسی_عرفان_کربالئی.docx'
d=Document(path)
for t in d.tables:
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                if p.text.replace('\u200e','').replace('\u00a0',' ')=='نوع SQL':
                    for r in p.runs:r.text=r.text.replace('SQL','ORM')
for p in d.paragraphs:
    if p.text.replace('\u200e','').replace('\u00a0',' ').startswith('جدول‌های زیر مستقیماً از metadata'):
        from build_report import mixed
        for r in p.runs:r._r.getparent().remove(r._r)
        mixed(p,'جدول‌های زیر مستقیماً از metadata مدل‌های SQLAlchemy استخراج شده‌اند. نوع نمایشی ORM، nullable، PK و FK از مدل برداشت شده‌اند. نوع عمومی DATETIME با نوع بومی PostgreSQL متفاوت است؛ ستون DateTime دارای timezone در PostgreSQL به timestamp with time zone نگاشت می‌شود و enum بومی نام نوع خود را دارد. مقادیر enum در جدول جدا آمده‌اند. هیچ مقدار واقعی بیمار یا حساب وجود ندارد. created_at و updated_at از TimestampMixin می‌آیند و Nullable با رشتهٔ خالی یکسان نیست.')
d.save(path)
print('ORM type labels clarified.')
