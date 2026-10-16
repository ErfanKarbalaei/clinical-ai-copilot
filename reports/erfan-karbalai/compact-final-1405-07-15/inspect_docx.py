from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
import re, sys
sys.stdout.reconfigure(encoding="utf-8")
root=Path(__file__).resolve().parents[1]
src=next(root.glob("*بازبینی‌شده.docx"))
doc=Document(src)
print("SOURCE",src,"PARAGRAPHS",len(doc.paragraphs),"TABLES",len(doc.tables),"IMAGES",len(doc.inline_shapes))
for i,p in enumerate(doc.paragraphs):
    s=p.text.strip()
    if p.style.name.startswith("Heading") or any(k in s for k in ("پیوست","منابع","عرفان کرب","۱۴۰۵","کربال")):
        print(f"P{i} [{p.style.name}] {s[:500]}")
print("\nBIBLIOGRAPHY CANDIDATES")
for i,p in enumerate(doc.paragraphs):
    s=p.text.strip()
    if 608 <= i <= 640:
        print(f"P{i}: {s[:1000]}")
print("\nAPPENDIX PARAGRAPHS")
for i,p in enumerate(doc.paragraphs):
    if 642 <= i < len(doc.paragraphs):
        s=p.text.strip()
        if s:
            print(f"P{i} [{p.style.name}]: {s[:320]}")
print("\nCITATION NUMBER OCCURRENCES")
alltext="\n".join(p.text for p in doc.paragraphs)+"\n"+"\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
for n in range(1,15):
    fa=str.maketrans("0123456789","۰۱۲۳۴۵۶۷۸۹")
    count=len(re.findall(r"(?<!\w)(?:\["+str(n)+r"\]|\["+str(n).translate(fa)+r"\]|【"+str(n)+r"】|【"+str(n).translate(fa)+r"】)",alltext))
    if count: print(n,count)
print("\nDOC BODY ORDER:")
for i,node in enumerate(doc.element.body):
    if node.tag==qn("w:p"):
        from docx.text.paragraph import Paragraph
        p=Paragraph(node,doc)
        s=p.text.strip()
        if p.style.name.startswith("Heading") or "عرفان کرب" in s or "پیوست" in s or "منابع" in s:
            print("BODY",i,"PARA",p.style.name,repr(s[:200]))
    elif node.tag==qn("w:tbl"):
        ix=sum(1 for n in list(doc.element.body)[:i+1] if n.tag==qn("w:tbl"))-1
        tab=doc.tables[ix]
        cells=[]
        for row in tab.rows[:2]:
            cells.append(" | ".join(c.text[:80].replace("\n"," / ") for c in row.cells))
        print("BODY",i,"TABLE",ix," :: ".join(cells))
