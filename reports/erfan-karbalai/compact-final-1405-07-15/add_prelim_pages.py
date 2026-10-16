from pathlib import Path
import shutil
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.shared import Pt

HERE = Path(__file__).resolve().parent
SOURCE = HERE / "گزارش_پروژه_کارشناسی_عرفان_کربلائی_خلاصه‌شده.docx"
OUTDIR = HERE.parent / "revised-1405-07-15"
OUTDIR.mkdir(parents=True, exist_ok=True)
OUTPUT = OUTDIR / "گزارش_پروژه_کارشناسی_عرفان_کربلائی_با_تقدیم_و_سپاسگزاری.docx"

doc = Document(SOURCE)
assert len(doc.paragraphs) > 20
assert doc.paragraphs[16].text.strip() == "صورت‌جلسه‌ی داوری"
assert doc.paragraphs[18]._p.xpath(".//w:br[@w:type='page']")
assert doc.paragraphs[19]._p.xpath(".//w:br[@w:type='page']")
anchor = doc.paragraphs[18]._p

def insert_para_before(anchor, text="", *, page_break=False, title=False):
    p_elm = OxmlElement("w:p")
    ppr = OxmlElement("w:pPr")
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    ppr.append(bidi)
    jc = OxmlElement("w:jc")
    jc.set(qn("w:val"), "start")
    ppr.append(jc)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:before"), "0" if page_break else ("300" if title else "180"))
    spacing.set(qn("w:after"), "240" if title else "0")
    spacing.set(qn("w:line"), "240")
    spacing.set(qn("w:lineRule"), "auto")
    ppr.append(spacing)
    p_elm.append(ppr)
    if page_break:
        run = OxmlElement("w:r")
        br = OxmlElement("w:br")
        br.set(qn("w:type"), "page")
        run.append(br)
        p_elm.append(run)
    elif text:
        run = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        rtl = OxmlElement("w:rtl")
        rtl.set(qn("w:val"), "1")
        rpr.append(rtl)
        lang = OxmlElement("w:lang")
        lang.set(qn("w:val"), "fa-IR")
        lang.set(qn("w:eastAsia"), "fa-IR")
        rpr.append(lang)
        fonts = OxmlElement("w:rFonts")
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn("w:" + attr), "B Nazanin")
        rpr.append(fonts)
        size = OxmlElement("w:sz")
        size.set(qn("w:val"), "40" if title else "28")
        rpr.append(size)
        if title:
            rpr.append(OxmlElement("w:b"))
        run.append(rpr)
        t = OxmlElement("w:t")
        t.text = text
        run.append(t)
        p_elm.append(run)
    anchor.addprevious(p_elm)
    return Paragraph(p_elm, doc._body)

# Add one clean Persian dedication page and one concise acknowledgements page.
insert_para_before(anchor, page_break=True)
insert_para_before(anchor, "تقدیم", title=True)
insert_para_before(anchor, "تقدیم به همه‌ی کسانی که دانش و فناوری را در خدمت به سلامت و کرامت انسان به کار می‌گیرند.")
insert_para_before(anchor, page_break=True)
insert_para_before(anchor, "سپاسگزاری", title=True)
insert_para_before(anchor, "از استاد راهنمای محترم پروژه، جناب آقای دکتر کاظم فولادی قلعه، و از دانشکده و گروه مهندسی کامپیوتر دانشگاه تهران سپاسگزاری می‌کنم.")

doc.core_properties.title = "گزارش پروژه کارشناسی عرفان کربلائی"
doc.core_properties.subject = "نسخه‌ی بازبینی‌شده با صفحه‌های تقدیم و سپاسگزاری"
doc.save(OUTPUT)
print("Saved", OUTPUT)
print("Paragraphs", len(doc.paragraphs), "tables", len(doc.tables), "images", len(doc.inline_shapes))
