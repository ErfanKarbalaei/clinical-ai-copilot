from pathlib import Path
from docx import Document
import re,sys
sys.stdout.reconfigure(encoding="utf-8")
root=Path(__file__).resolve().parents[1]
doc=Document(next(root.glob("*بازبینی‌شده.docx")))
print("BIBLIOGRAPHY")
for i,p in enumerate(doc.paragraphs):
 s=p.text.strip()
 if i>=608 and i<642 and s: print(i,repr(s))
print("ENDAPP")
for i,p in enumerate(doc.paragraphs):
 s=p.text.strip()
 if i>=850 and s: print(i,p.style.name,repr(s[:200]))
alltext="\n".join(p.text for p in doc.paragraphs)+"\n"+"\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
print("CITATIONS")
for n in range(1,15):
 fa=str.maketrans("0123456789","۰۱۲۳۴۵۶۷۸۹")
 pat=r"(?:\["+str(n)+r"\]|\["+str(n).translate(fa)+r"\])"
 contexts=[m.start() for m in re.finditer(pat,alltext)]
 if contexts: print(n,len(contexts))
