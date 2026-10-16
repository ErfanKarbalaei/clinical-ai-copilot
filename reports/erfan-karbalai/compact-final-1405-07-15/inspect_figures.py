from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
import sys
sys.stdout.reconfigure(encoding="utf-8")
root=Path(__file__).resolve().parents[1];d=Document(next(root.glob("*بازبینی‌شده.docx")))
for i,p in enumerate(d.paragraphs):
    blips=p._p.xpath(".//a:blip")
    if blips:
        captions=[]
        for q in d.paragraphs[i+1:i+4]:
            if q.text.strip():
                captions.append(q.text.strip())
                if q.text.strip().startswith("شکل"): break
        for b in blips:
            rid=b.get(qn("r:embed"));part=d.part.related_parts[rid]
            print(i,part.partname,repr(p.text[:50]),"NEXT",repr(captions))
for i,p in enumerate(d.paragraphs[10:20],10):
    print("PRE",i,repr(p.text),p._p.xml[-700:])
