from pathlib import Path
from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn
import re,sys
sys.stdout.reconfigure(encoding="utf-8")
root=Path(__file__).resolve().parents[1];d=Document(next(root.glob("*بازبینی‌شده.docx")))
print("PRELIM")
for i,p in enumerate(d.paragraphs[:245]):
 if p.text.strip() or p._p.xpath(".//w:drawing"):
  print(i,p.style.name,repr(p.text[:200]),"draw",len(p._p.xpath(".//w:drawing")),"pageBreak",p.paragraph_format.page_break_before)
print("FIGURES")
for i,p in enumerate(d.paragraphs):
 if p._p.xpath(".//w:drawing"):
  before=[d.paragraphs[j].text.strip() for j in range(max(0,i-2),i) if d.paragraphs[j].text.strip()]
  after=[d.paragraphs[j].text.strip() for j in range(i+1,min(len(d.paragraphs),i+3)) if d.paragraphs[j].text.strip()]
  print(i,"style",p.style.name,"text",repr(p.text[:100]),"before",before,"after",after)
print("CITES STRIPPED")
alltext="\n".join(p.text for p in d.paragraphs)+"\n"+"\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)
alltext=alltext.replace("\u200e","").replace("\u200f","").replace("\u2066","").replace("\u2069","")
for m in re.finditer(r"\[\s*([0-9۰-۹]{1,2})\s*\]",alltext):
 print(m.group(1),repr(alltext[max(0,m.start()-65):m.end()+65]))
