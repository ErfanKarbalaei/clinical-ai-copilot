"""Read-only structural cross-check of API, ORM, configuration and excerpts."""
from pathlib import Path
import sys,json,ast,hashlib,re
from docx import Document
OUT=Path(__file__).resolve().parent;ROOT=OUT.parents[1]
sys.stdout.reconfigure(encoding='utf-8');sys.path.insert(0,str(ROOT/'backend'))
from app.main import app as application
from app.db.base import Base
import app.models
from app.core.config import Settings
from app.services import ai,stt,patient_context,evidence
def clean(t):return t.replace('\u200e','').replace('\u00a0',' ')
d=Document(OUT/('گزارش_پروژه_کارشناسی_عرفان_کربالئی_بازبینی‌شده.docx' if '--revised' in sys.argv else 'گزارش_پروژه_کارشناسی_عرفان_کربالئی.docx'))
routes={(method.upper(),path) for path,ops in application.openapi()['paths'].items() for method in ops if method.upper() in ['GET','POST','PATCH','PUT','DELETE']}
reported=[]
for t in d.tables:
    for row in t.rows:
        cells=[clean(c.text).strip() for c in row.cells]
        for method in ['GET','POST','PATCH','PUT','DELETE']:
            if method in cells:
                p=next((c for c in cells if c.startswith('/')),None)
                if p:reported.append((method,p))
        combined=next((re.match(r'^(GET|POST|PATCH|PUT|DELETE)\s+(/[\s\S]+)$',c) for c in cells if re.match(r'^(GET|POST|PATCH|PUT|DELETE)\s+(/[\s\S]+)$',c)),None)
        if combined:reported.append((combined.group(1),''.join(combined.group(2).split())))
api_missing=sorted(set(reported)-routes);unreported=sorted(routes-set(reported))
orm=[]
for idx,t in enumerate(d.tables):
    headers=[clean(c.text) for c in t.rows[0].cells]
    if 'نوع ORM' in headers:
        names={clean(row.cells[0].text).strip() for row in t.rows[1:]}
        matches=[table.name for table in Base.metadata.sorted_tables if names==set(table.columns.keys())]
        orm.append({'table_index':idx,'fields':sorted(names),'matching_orm_tables':matches})
tree=ast.parse((ROOT/'backend/app/api/v1/api.py').read_text(encoding='utf-8'))
routers=[n.args[0].id for n in ast.walk(tree) if isinstance(n,ast.Call) and isinstance(n.func,ast.Attribute) and n.func.attr=='include_router']
docs=(ROOT/'backend/app/services/document_analysis.py').read_text(encoding='utf-8')
tool=(ROOT/'backend/app/mcp_server.py').read_text(encoding='utf-8')
code=[p.text for p in d.paragraphs if p.style.name=='Code']
enum_migration_files=sorted(p.name for p in (ROOT/'backend/alembic/versions').glob('*.py'))
source_paths=list((ROOT/'backend/app').rglob('*.py'))+list((ROOT/'backend/tests').glob('*.py'))+list((ROOT/'backend/alembic/versions').glob('*.py'))
source_paths += [ROOT/p for p in ['frontend/src/lib/api.ts','frontend/src/lib/auth.ts','frontend/src/components/patients/clinical-chat.tsx','frontend/src/components/patients/medication-actions.tsx','frontend/src/components/patients/copilot/evidence-list.tsx','backend/Dockerfile','frontend/Dockerfile','deploy/docker-compose.prod.yml','deploy/nginx/nginx.conf']]
hashes={str(p.relative_to(ROOT)).replace('\\','/'):hashlib.sha256(p.read_bytes()).hexdigest() for p in source_paths}
result={'router_count':len(routers),'registered_routers':routers,'api_count':len(routes),'reported_api_count':len(reported),'unsupported_reported_routes':api_missing,'unreported_routes':unreported,'orm_field_tables':orm,'setting_field_names':sorted(Settings.model_fields),'migration_count':len(enum_migration_files),'migrations':enum_migration_files,'context_limit':patient_context._RECENT_LIMIT,'evidence_limit':evidence.MAX_EVIDENCE_ITEMS,'text_providers':list(ai.PROVIDER_IDS),'stt_classes':[name for name in ['SpeechmaticsProvider','GoogleSpeechProvider','OpenRouterSTTProvider'] if hasattr(stt,name)],'mcp_tool_count':tool.count('server.add_tool('),'code_blocks_count':len(code),'code_hashes':[hashlib.sha256(c.encode()).hexdigest() for c in code],'source_hashes':hashes}
(OUT/('review-technical-final.json' if '--revised' in sys.argv else 'review-technical-check.json')).write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:result[k] for k in ['router_count','api_count','reported_api_count','unsupported_reported_routes','unreported_routes','migration_count','context_limit','evidence_limit','mcp_tool_count']},ensure_ascii=False,indent=2))
print('ORM field tables matched:',sum(bool(r['matching_orm_tables']) for r in orm),'/',len(orm))
