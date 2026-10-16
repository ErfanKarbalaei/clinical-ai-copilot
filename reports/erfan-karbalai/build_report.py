"""Build the editable Persian report from inspected content and real metadata.
All outputs remain in this directory. Does not connect to application services.
"""
from pathlib import Path
import sys,os,re,json,subprocess,ast,hashlib,xml.etree.ElementTree as ET
OUT=Path(__file__).resolve().parent
ROOT=OUT.parents[1]
sys.path.insert(0,str(ROOT/'backend'))
sys.stdout.reconfigure(encoding='utf-8')
os.environ['PYTHONDONTWRITEBYTECODE']='1'
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_BREAK
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image,ImageDraw,ImageFont
import pymupdf

TITLE='طراحی و پیاده‌سازی دستیار هوشمند بالینی (Clinical AI Copilot) مبتنی بر تلفیق پرونده بیمار و دانش پزشکی معتبر'
FA='B Nazanin'; EN='Times New Roman'
DOCX=OUT/'گزارش_پروژه_کارشناسی_عرفان_کربالئی.docx'
PERSIAN_DIGITS=str.maketrans('0123456789','۰۱۲۳۴۵۶۷۸۹')
def fa_num(x):return str(x).translate(PERSIAN_DIGITS)
def xml(parent,tag,**attrs):
    el=OxmlElement('w:'+tag)
    for k,v in attrs.items():el.set(qn('w:'+k),str(v))
    parent.append(el);return el
def set_bidi(p,rtl=True):
    pr=p._p.get_or_add_pPr()
    for n in list(pr.findall(qn('w:bidi'))):pr.remove(n)
    xml(pr,'bidi',val='1' if rtl else '0')
def run_format(r,size=14,rtl=True,bold=False,font=None):
    r.font.name=font or (FA if rtl else EN);r.font.size=Pt(size)
    r.font.bold=bold;r.font.color.rgb=RGBColor(0,0,0)
    pr=r._r.get_or_add_rPr()
    rf=pr.find(qn('w:rFonts'))
    if rf is None:rf=xml(pr,'rFonts')
    for k in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+k),font or (FA if rtl else EN))
    xml(pr,'rtl',val=1 if rtl else 0);xml(pr,'szCs',val=int(size*2))
    xml(pr,'bCs',val=1 if bold else 0);xml(pr,'lang',val='fa-IR' if rtl else 'en-US',bidi='fa-IR')
    return r
# Keep a whole path, identifier, URL, or multiword English name in one LTR run.
TOKEN=re.compile(r'(?:[۰-۹]+(?:ـ[۰-۹]+)+|https?://[^\s]+|/?[A-Za-z][A-Za-z0-9_./:{}\[\]?=+%<>#-]*(?: [A-Za-z][A-Za-z0-9_./:{}\[\]?=+%<>#-]*)*|\d+(?:[.:/][\d]+)*|_[A-Za-z][A-Za-z0-9_]*)')
def mixed(p,text,size=14,bold=False):
    text=text.replace('هٔ','ه‌ی')
    pos=0
    for m in TOKEN.finditer(text):
        if m.start()>pos:run_format(p.add_run(text[pos:m.start()]),size,True,bold)
        # LRM inside explicit LTR run stabilizes punctuation at RTL boundaries.
        token=m.group()
        if re.fullmatch('[۰-۹]+(?:ـ[۰-۹]+)+',token):token=token.replace('ـ','-')
        if len(token)<38:token=token.replace(' ','\u00a0')
        run_format(p.add_run('\u200e'+token+'\u200e'),max(size-1,10),False,bold)
        pos=m.end()
    if pos<len(text):run_format(p.add_run(text[pos:]),size,True,bold)
def para(text='',style='Normal',size=14,align=None,bold=False):
    p=doc.add_paragraph(style=style);set_bidi(p)
    p.paragraph_format.alignment=align if align is not None else WD_ALIGN_PARAGRAPH.JUSTIFY
    if align==WD_ALIGN_PARAGRAPH.RIGHT:
        jc=p._p.get_or_add_pPr().find(qn('w:jc'));jc.set(qn('w:val'),'start')
    mixed(p,text,size,bold);return p
def ltr(text,size=10):
    p=doc.add_paragraph(style='Code');set_bidi(p,False)
    p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    run_format(p.add_run(text),size,False,font='Consolas');return p
def field(p,instr):
    xml(p.add_run()._r,'fldChar',fldCharType='begin')
    t=xml(p.add_run()._r,'instrText');t.set(qn('xml:space'),'preserve');t.text=instr
    xml(p.add_run()._r,'fldChar',fldCharType='separate')
    run_format(p.add_run('۱'),11)
    xml(p.add_run()._r,'fldChar',fldCharType='end')
def bookmark(p,name):
    global BOOKMARK
    BOOKMARK+=1
    st=OxmlElement('w:bookmarkStart');st.set(qn('w:id'),str(BOOKMARK));st.set(qn('w:name'),name)
    p._p.insert(0,st);en=OxmlElement('w:bookmarkEnd');en.set(qn('w:id'),str(BOOKMARK));p._p.append(en)
    return name
BOOKMARK=0; chapter_records=[]; table_idx=0; figures=[];tables=[]
def margins(s):
    s.page_width=Cm(21);s.page_height=Cm(29.7)
    s.top_margin=Cm(3.5);s.bottom_margin=Cm(2)
    s.left_margin=Cm(2);s.right_margin=Cm(2.5)
    s.header_distance=Cm(2.5);s.footer_distance=Cm(1)
def clear(h):
    h.is_linked_to_previous=False
    for p in h.paragraphs:
        for ch in list(p._p):
            if ch.tag!=qn('w:pPr'):p._p.remove(ch)
def setup_headers(s,title=None,first=False):
    s.different_first_page_header_footer=first
    for part in [s.header,s.even_page_header,s.first_page_header,s.footer,s.even_page_footer,s.first_page_footer]:clear(part)
    if title:
        for even,h in [(False,s.header),(True,s.even_page_header)]:
            base=h.paragraphs[0];base.paragraph_format.space_after=Pt(0);base.paragraph_format.space_before=Pt(0)
            base.paragraph_format.line_spacing=Pt(1)
            run_format(base.add_run(''),1)
            t=h.add_table(rows=1,cols=2,width=Cm(16.5));t.autofit=False
            xml(t._tbl.tblPr,'bidiVisual',val=1)
            for col in t.columns:col.width=Cm(8.25)
            for ci,cell in enumerate(t.rows[0].cells):
                cell.width=Cm(8.25)
                p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0)
                isnum=(ci==0) if even else (ci==1)
                if isnum:
                    set_bidi(p,False);p.alignment=WD_ALIGN_PARAGRAPH.RIGHT if even else WD_ALIGN_PARAGRAPH.LEFT
                    field(p,' PAGE ')
                else:
                    set_bidi(p);xml(p._p.get_or_add_pPr(),'jc',val='end' if even else 'start');mixed(p,title,10)
                borders=xml(cell._tc.get_or_add_tcPr(),'tcBorders');xml(borders,'bottom',val='single',sz=4,color='000000')
        if first:
            p=s.first_page_footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;field(p,' PAGE ')
def section(title=None,first=True,start=None):
    s=doc.add_section(WD_SECTION_START.ODD_PAGE);margins(s);setup_headers(s,title,first)
    bp=doc.paragraphs[-1]
    bp.paragraph_format.line_spacing=Pt(1)
    bp.paragraph_format.space_before=Pt(0);bp.paragraph_format.space_after=Pt(0)
    bp.paragraph_format.keep_with_next=False
    for n in list(s._sectPr.findall(qn('w:pgNumType'))):s._sectPr.remove(n)
    if start is not None:xml(s._sectPr,'pgNumType',start=start)
    return s
def chapter(title,kind='main'):
    section(title)
    p=para(title,'Heading 1',20,WD_ALIGN_PARAGRAPH.RIGHT,True)
    name=bookmark(p,f'section_{len(chapter_records)+1}')
    chapter_records.append({'title':title,'kind':kind,'bookmark':name})
def heading(text,level=2):return para(text,'Heading '+str(level),18 if level==2 else 16,WD_ALIGN_PARAGRAPH.RIGHT,True)
def caption(text,figure=False):
    p=para(text,'ReportFigure' if figure else 'ReportTable',12,WD_ALIGN_PARAGRAPH.CENTER)
    p.paragraph_format.keep_with_next=not figure
    (figures if figure else tables).append(text)
    return p
def table(title,headers,rows,widths=None):
    global table_idx
    table_idx+=1;caption('جدول '+fa_num(table_idx)+' '+title)
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
    xml(t._tbl.tblPr,'bidiVisual',val=1)
    widths=widths or [16.5/len(headers)]*len(headers)
    for col,w in zip(t.columns,widths):col.width=Cm(w)
    for ridx,data in enumerate([headers]+list(rows)):
        cells=t.rows[0].cells if ridx==0 else t.add_row().cells
        pr=t.rows[ridx]._tr.get_or_add_trPr();xml(pr,'cantSplit')
        if ridx==0:xml(pr,'tblHeader')
        for ci,(cell,text) in enumerate(zip(cells,data)):
            cell.width=Cm(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc=cell._tc.get_or_add_tcPr()
            if ridx==0:xml(tc,'shd',fill='E7E7E7',val='clear')
            borders=xml(tc,'tcBorders')
            for side in ['top','left','bottom','right']:xml(borders,side,val='single',sz=4,color='D9D9D9')
            pad=xml(tc,'tcMar')
            for side in ['top','bottom']:xml(pad,side,w=65,type='dxa')
            for side in ['left','right']:xml(pad,side,w=90,type='dxa')
            p=cell.paragraphs[0];p.style='ReportCell';set_bidi(p)
            if ridx==0:p.paragraph_format.keep_with_next=True
            xml(p._p.get_or_add_pPr(),'jc',val='start')
            # A pure technical field is LTR and left-aligned. Persian prose is RTL.
            if not re.search('[\u0600-\u06ff]',str(text)):
                set_bidi(p,False);p.alignment=WD_ALIGN_PARAGRAPH.LEFT
                run_format(p.add_run(str(text)),10.5,False,ridx==0)
            else:mixed(p,str(text),12,ridx==0)
    para('',size=8).paragraph_format.space_after=Pt(0)
    return t

def draw_diagram(name,nodes,edges,size=(1600,800),routes=None):
    """Code-drawn technical figures, English labels and Persian DOCX captions."""
    im=Image.new('RGB',size,'white');dr=ImageDraw.Draw(im)
    font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',28)
    small=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',23)
    paths=list(routes or [])
    for a,b in edges:
        x,y,w,h,_=nodes[a];xx,yy,ww,hh,_=nodes[b]
        if abs(xx-x)>abs(yy-y):
            start=(x+w,y+h/2) if xx>x else (x,y+h/2)
            end=(xx,yy+hh/2) if xx>x else (xx+ww,yy+hh/2)
        else:
            start=(x+w/2,y+h) if yy>y else (x+w/2,y)
            end=(xx+ww/2,yy) if yy>y else (xx+ww/2,yy+hh)
        paths.append([start,end])
    for points in paths:
        dr.line(points,fill='#344454',width=4)
        start,end=points[-2:]
        import math
        ang=math.atan2(end[1]-start[1],end[0]-start[0]);r=17
        dr.polygon([end,(end[0]-r*math.cos(ang-.45),end[1]-r*math.sin(ang-.45)),(end[0]-r*math.cos(ang+.45),end[1]-r*math.sin(ang+.45))],fill='#344454')
    for x,y,w,h,label in nodes:
        dr.rounded_rectangle((x,y,x+w,y+h),radius=12,fill='#F4F6F8',outline='#344454',width=3)
        lines=label.split('\n');dy=(h-len(lines)*35)/2
        for j,line in enumerate(lines):
            f=font if j==0 else small
            box=dr.textbbox((0,0),line,font=f);dr.text((x+(w-(box[2]-box[0]))/2,y+dy+j*35),line,font=f,fill='black')
    path=OUT/'figures'/f'{name}.png';path.parent.mkdir(exist_ok=True);im.save(path)
def flow(name,labels,extra=None):
    nodes=[]
    for i,label in enumerate(labels):
        row=i//3;col=i%3 if row%2==0 else 2-i%3
        nodes.append((45+col*535,50+row*230,460,150,label))
    draw_diagram(name,nodes,list(zip(range(len(labels)-1),range(1,len(labels)))),(1600, max(540,((len(labels)+2)//3)*230+60)))
def diagrams():
    draw_diagram('architecture',[(70,70,420,140,'Next.js\nClinical workflow UI'),(600,70,420,140,'FastAPI\nRoutes / access / services'),(1110,70,420,140,'External providers\nLLM / STT APIs'),(600,340,420,140,'PostgreSQL\nRecords / messages / evidence'),(1110,340,420,140,'MinIO\nOriginal documents / audio'),(70,340,420,140,'Nginx\nProduction HTTP entry')],[(0,1),(1,2),(1,3),(1,4),(5,0),(5,1)],(1600,550))
    flow('request',['HTTP request\nJSON or multipart','Pydantic\nInput validation','Dependencies\nUser + patient access','Route / service\nBusiness operation','DB / provider\nCommit or failure','Response schema\nStatus + JSON'])
    draw_diagram('erd',[(40,20,460,120,'users\nAuthentication'),(40,230,460,120,'patient_user_access\nUnique patient / user'),(600,20,460,120,'patients\nClinical record root'),(40,500,460,140,'conversations\nconversation_messages'),(575,500,460,140,'patient_medications\nmedications catalog'),(1110,500,460,140,'documents / voice_notes\nFile keys + derived text'),(1110,750,460,140,'lab_results / clinical_notes\nOptional document / voice FK'),(40,750,460,140,'conditions / allergies\nPatient-scoped data')],[],(1600,940),routes=[[(270,140),(270,230)],[(600,80),(550,80),(550,290),(500,290)],[(830,140),(830,430),(270,430),(270,500)],[(830,140),(830,430),(805,430),(805,500)],[(830,140),(830,430),(1340,430),(1340,500)],[(830,140),(830,400),(20,400),(20,820),(40,820)],[(830,140),(830,400),(1590,400),(1590,820),(1570,820)],[(1340,640),(1340,750)]])
    flow('context',['Authorized patient\nReal database ID','Seven queries\nStatus / date ordering','Per-section cap\n200 records','PatientContextResponse\nStructured live data','Consumers\nUI / chat / safety / MCP'])
    flow('documents',['Multipart upload\nMIME / 20 MiB limit','MinIO + Document\nOriginal file + metadata','BackgroundTasks\nIndependent DB session','Text / PDF / OCR\nExtracted text','Optional AI analysis\nSeparate ai_summary','Authorized download\n307 presigned URL'])
    flow('chat',['Conversation access\nHistory: last 20 messages','Commit user message\nQuestion survives AI failure','Live patient context\nReal evidence IDs','Provider completion\nPrompt + context + history','Parse + validate IDs\nDiscard invalid references','Commit assistant\nAnswer + JSONB evidence'])
    draw_diagram('stt',[(35,35,465,140,'VoiceNote upload\nOriginal audio in MinIO'),(575,35,465,140,'transcribe_audio\nConfigured primary provider'),(1110,35,465,140,'Fallback chain\nOnly after exceptions'),(35,285,465,150,'Speechmatics\nBatch jobs / speaker tags'),(575,285,465,150,'Google web Speech\n55s chunks / retries'),(1110,285,465,150,'OpenRouter audio\nFull WAV / input_audio'),(575,570,465,150,'Raw transcript\nOptional cleanup / reviewed note')],[(0,1),(1,2),(1,3),(1,4),(1,5),(3,6),(4,6),(5,6)],(1600,770))
    draw_diagram('safety',[(40,60,450,150,'Candidate medication\nPhysician + patient context'),(575,40,450,150,'AI-first workflow\nJSON analysis / chat mode'),(1100,40,450,150,'Provider model\nGenerated suggestions'),(575,290,450,150,'Legacy rule endpoint\nLocal limited rules'),(1100,290,450,150,'Deterministic warnings\nRule ID / coverage'),(40,530,450,150,'Save medication\nExplicit physician action')],[(0,1),(1,2),(0,3),(3,4),(0,5)],(1600,730))
    flow('evidence',['Real context records\nIDs from primary keys','Annotate context\nDOC / LAB / MED / etc.','Model JSON response\nanswer + evidence IDs','validate_evidence\nPatient index / cap 10','Persist metadata\nNo temporary URL','Source access\nRecord tab / document endpoint'])
    flow('auth',['Bearer token\nJWT signature + expiry','Active User\nDatabase check','Patient relationship\nAdmin or explicit grant','Child source patient_id\nNo client-supplied ownership','Authorized operation\nDB / AI / storage'])
    draw_diagram('deployment',[(40,60,450,140,'Nginx\nOnly published entry'),(575,60,450,140,'Frontend :3000\nNext.js standalone'),(1100,60,450,140,'Backend :8000\nFastAPI / Uvicorn'),(575,350,450,140,'MinIO :9000\nObject volume'),(1100,350,450,140,'PostgreSQL :5432\nDatabase volume'),(40,350,450,140,'Storage hostname\nHost preserved for signature')],[(0,1),(0,2),(2,3),(2,4),(0,5),(5,3)],(1600,560))

DATA={
'responsibilities':('تقسیم مسئولیت مطابق پروپوزال',['حوزه','مسئول اصلی','مرز مسئولیت'],[
['Backend و AI','عرفان کربالئی','API و داده، سند و صوت، Context، اتصال مدل، شواهد و ایمنی'],['Frontend و UX','ایمان کربلائی','رابط بالینی، نمایش پرونده و خروجی، ضبط و بارگذاری و اتصال UI'],['معماری و تحویل','هر دو دانشجو','قرارداد API، Integration، E2E، Deployment و Demo']]),
'requirements':('نیاز و سطح تحقق',['نیاز','پیاده‌سازی موجود','وضعیت ارزیابی'],[
['مدیریت بیمار','CRUD و انتساب پزشک','کد و آزمون API موجود؛ اجرای DB در این گزارش انجام نشد'],['مرور پرونده','overview و Context و timeline','تجمیع قطعی؛ بدون ارزیابی زمان مرور'],['پردازش سند','متن PDF/OCR و ai_summary','استخراج ساختاریافتهٔ فیلد خودکار نیست'],['مستندسازی صوت','STT و متن خام و یادداشت بازبینی‌شده','providerها موجود؛ WER/CER اندازه‌گیری نشده'],['گفتگو و ارجاع','پیام ماندگار و evidence واقعی','اعتبار ID کنترل می‌شود؛ صحت معنایی جداست'],['ایمنی نسخه','AI-first و موتور قدیمی مستقل','فاقد اعتبارسنجی بالینی'],['دانش بیرونی','در طرح اختیاری/آینده','PubMed و RxNorm و FDA متصل نشده‌اند']]),
'timeline':('منبع و تاریخ رویدادهای تایم‌لاین',['منبع','رویداد','تاریخ اصلی و fallback'],[
['Condition','condition_onset / condition_resolved','onset_date یا resolved_date؛ نبود تاریخ شروع → created_at'],['PatientMedication','medication_started / medication_stopped','start_date؛ پایان فقط با end_date و وضعیت غیرactive'],['Document','document_uploaded','created_at'],['VoiceNote','voice_note_recorded','created_at'],['ClinicalNote','note_created','created_at'],['LabResult','lab_result','result_date؛ در نبود → created_at']]),
'ai_providers':('قرارداد ارائه‌دهنده‌های متن',['ارائه‌دهنده','قالب تماس','کلید و محدودیت'],[
['OpenAICompatibleProvider','chat/completions و messages','AI_API_KEY؛ base URL از سرور'],['GeminiProvider','generateContent و system_instruction','AI_API_KEY؛ roles به مدل نگاشت می‌شوند'],['OpenRouterProvider','chat/completions و X-Title','OPENROUTER_API_KEY سپس AI_API_KEY؛ نام مدل تضمین availability نیست']]),
'objective_status':('تطبیق خروجی با هدف‌های طرح',['هدف','وضعیت نسخهٔ موجود','کار باقی‌مانده'],[
['هستهٔ داده و API','موجود با مدل و migration','تقویت تراکنش و سیاست حذف فایل'],['اسناد و مقادیر منتخب','متن و تحلیل سند موجود','فیلدهای معتبر و provenance جزئی'],['Context و timeline','موجود و زنده','انتخاب مرتبط و اعلام truncation'],['مدل آماده و گفتگو','سه adapter متن و گفتگو','ارزیابی معنایی و بازتولید پاسخ'],['Citation','رکورد پرونده','پشتیبانی معنایی و بازیابی بیرونی'],['Speech-to-Text','سه مسیر provider','ارزیابی واقعی و metadata نتیجه'],['Prescription Safety','AI-first و قواعد مستقل','تطبیق ایمنی با طرح و مرجع معتبر'],['استقرار','Docker و Nginx پیکربندی شده','TLS و اجرای stack و backup'],['ارزیابی کاربر و بالینی','نتیجهٔ تأییدشده موجود نیست','پروتکل دادهٔ مرجع و مرور انسانی']])
}
def main_tables(key):
    if key=='migrations':
        rows=[]
        for f in (ROOT/'backend/alembic/versions').glob('*.py'):
            tree=ast.parse(f.read_text(encoding='utf-8'));info={}
            for n in tree.body:
                if isinstance(n,ast.AnnAssign) and isinstance(n.target,ast.Name) and n.target.id in ['revision','down_revision']:
                    info[n.target.id]=ast.literal_eval(n.value)
            rows.append([info.get('revision'),info.get('down_revision') or 'آغاز',f.stem.split('_',1)[1]])
        ordered=[];parent=None
        while rows:
            m=next((r for r in rows if r[1]==(parent or 'آغاز')),None)
            if m is None:break
            ordered.append(m);rows.remove(m);parent=m[0]
        table('زنجیرهٔ واقعی مهاجرت‌ها',['Revision','پیشین','تغییر'],ordered,[3.3,3.3,9.9])
    elif key=='test_results':
        rows=[]
        tree=ET.parse(OUT/'unit-tests.xml')
        grouped={}
        for case in tree.findall('.//testcase'):
            group=case.get('classname','').split('.')[-1]
            v=grouped.setdefault(group,[0,0,0]);v[1 if case.find('error') is not None or case.find('failure') is not None else 2 if case.find('skipped') is not None else 0]+=1
        for group,c in grouped.items():rows.append([group,str(c[0]),str(c[1]),str(c[2])])
        table('نتیجهٔ اجرای منتخب',['فایل آزمون','موفق','خطا','کنارگذاشته'],rows,[10,2,2,2.5])
    else:
        title,headers,rows=DATA[key];table(title,headers,rows,[4,4.5,8])

def setup_styles():
    for name in ['ReportCell','ReportFigure','ReportTable','Code']:
        if name not in doc.styles:doc.styles.add_style(name,1)
    for name in ['Normal','Heading 1','Heading 2','Heading 3','Title','ReportCell','ReportFigure','ReportTable','Code']:
        st=doc.styles[name];st.font.name=FA;st.font.color.rgb=RGBColor(0,0,0)
        for b in st.element.xpath('./w:pPr/w:pBdr'):b.getparent().remove(b)
        sizes={'Heading 1':20,'Heading 2':18,'Heading 3':16,'Title':20,'ReportCell':12,'ReportFigure':12,'ReportTable':12,'Code':10}
        st.font.size=Pt(sizes.get(name,14));pr=st.element.get_or_add_rPr()
        xml(pr,'rFonts',ascii=EN,hAnsi=EN,cs=FA);xml(pr,'szCs',val=sizes.get(name,14)*2);xml(pr,'lang',val='fa-IR',bidi='fa-IR')
        pf=st.paragraph_format;pf.line_spacing=1.0;pf.space_after=Pt(5 if name=='Normal' else 3)
        pf.widow_control=True
        if name.startswith('Heading'):
            pf.keep_with_next=True;pf.space_before=Pt(10);pf.space_after=Pt(6)
        if name=='Normal':pf.first_line_indent=Cm(.5)
        if name=='ReportCell':pf.space_after=Pt(0);pf.first_line_indent=Cm(0)
        if name=='Code':pf.first_line_indent=Cm(0);pf.space_after=Pt(5)
    xml(doc.settings.element,'mirrorMargins')
    xml(doc.settings.element,'evenAndOddHeaders')
    xml(doc.settings.element,'updateFields',val=1)
    xml(doc.settings.element,'themeFontLang',val='fa-IR',bidi='fa-IR')

def preliminaries():
    s=doc.sections[0];margins(s);setup_headers(s)
    pdf=pymupdf.open(ROOT/'report-inputs/BSc-project-report-cover-fa-1401.pdf')
    pdf[0].get_pixmap(matrix=pymupdf.Matrix(3,3),clip=pymupdf.Rect(252,52,346,143)).save(str(OUT/'university-logo.png'))
    p=para('',align=WD_ALIGN_PARAGRAPH.CENTER);p.add_run().add_picture(str(OUT/'university-logo.png'),width=Cm(2.7))
    for t in ['دانشگاه تهران','دانشکدگان فارابی','دانشکدهٔ مهندسی','گروه مهندسی کامپیوتر']:para(t,size=16,align=WD_ALIGN_PARAGRAPH.CENTER,bold=True)
    p=para('طراحی و پیاده‌سازی دستیار هوشمند بالینی\n','Title',20,WD_ALIGN_PARAGRAPH.CENTER,True)
    run_format(p.add_run('(Clinical AI Copilot)\n'),19,False,True)
    mixed(p,'مبتنی بر تلفیق پرونده بیمار و دانش پزشکی معتبر',20,True)
    p.paragraph_format.space_before=Pt(35);p.paragraph_format.space_after=Pt(30)
    for t,b in [('نگارش',False),('عرفان کربالئی',True),('استاد راهنما',False),('دکتر کاظم فولادی قلعه',True)]:para(t,size=16,align=WD_ALIGN_PARAGRAPH.CENTER,bold=b)
    p=para('گزارش پروژه برای دریافت درجهٔ کارشناسی\nدر رشتهٔ مهندسی کامپیوتر',size=16,align=WD_ALIGN_PARAGRAPH.CENTER,bold=True);p.paragraph_format.space_before=Pt(35)
    para('ماه و سال تحویل: ................',size=14,align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_page_break() # title back page
    doc.add_page_break() # page 3 Bismillah kept completely empty
    doc.add_page_break() # page 4 blank back
    doc.add_page_break() # page 5 defense placeholder
    p=para('صورت‌جلسهٔ داوری',size=18,align=WD_ALIGN_PARAGRAPH.CENTER,bold=True);p.paragraph_format.space_before=Pt(70)
    para('محل درج فرم رسمی داوری پس از دریافت از دانشگاه',align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_page_break();doc.add_page_break() # page 6 blank -> page 7 abstract
    para('چکیده',size=18,align=WD_ALIGN_PARAGRAPH.RIGHT,bold=True)
    para('این پروژه یک نمونهٔ دانشگاهی تحت وب برای ثبت و مرور داده‌های بیمار و استفاده از ابزارهای آمادهٔ هوش مصنوعی است. هستهٔ سمت سرور با FastAPI و PostgreSQL اطلاعات بالینی، اسناد، یادداشت‌های صوتی و گفتگو را به بیمار مشخص مرتبط می‌کند. فایل اولیه در MinIO نگهداری می‌شود و متن استخراج‌شده، تحلیل سند و متن پاک‌سازی‌شدهٔ صوت از اصل منبع جدا می‌مانند. زمینهٔ بیمار از رکوردهای جاری ساخته می‌شود و تایم‌لاین ترتیب رخدادهای دارای منبع را عرضه می‌کند.')
    para('تمرکز گزارش بر مسئولیت‌های Backend و AI عرفان کربالئی است؛ مسئولیت اصلی رابط طبق پروپوزال با ایمان کربلائی و معماری و یکپارچه‌سازی مشترک است. اتصال مدل متن از سه adapter، گفتگوی ماندگار، کنترل شناسهٔ شواهد، سه مسیر گفتار به متن، احراز هویت و اجازهٔ پرونده در کد وجود دارند. مسیر فعلی تحلیل ایمنی دارو بر مدل تکیه دارد و از موتور محدود قواعد قدیمی مستقل است؛ این تفاوت با طرح اولیه به‌طور صریح بررسی می‌شود.')
    para('نسخهٔ موجود استخراج ساختاریافتهٔ خودکار مقادیر آزمایش، بازیابی بیرونی مقاله، اتصال عملیاتی بیمارستانی یا ارزیابی بالینی تأییدشده ندارد. در اجرای منتخب تازه، ۵۹ آزمون موفق و دو آزمون به دلیل آماده نبودن MinIO دچار خطای setup شدند. این نتیجه دربارهٔ صحت قراردادهای آزموده‌شده است و به کیفیت پزشکی مدل تعمیم داده نمی‌شود. دستاورد پروژه، مبنایی قابل بررسی برای یکپارچه‌سازی داده و سرویس است؛ کارهای آینده بر امنیت عملیاتی، وضعیت پردازش پایدار، منشأ دقیق و ارزیابی مبتنی بر دادهٔ مرجع متمرکزند.')
    para('کلیدواژه‌ها: دستیار هوشمند بالینی، زمینهٔ بیمار، FastAPI، پردازش سند، گفتار به متن، شواهد پرونده، ایمنی دارو.',bold=True)
    doc.add_page_break() # blank 8 -> section break creates odd TOC on page 9
    section('فهرست‌ها',False,start=1)
    para('فهرست مطالب',size=20,bold=True,align=WD_ALIGN_PARAGRAPH.RIGHT)
    p=para();field(p,' TOC \\o "1-2" \\h \\z \\u ')
    doc.add_page_break();para('فهرست شکل‌ها',size=18,bold=True,align=WD_ALIGN_PARAGRAPH.RIGHT)
    p=para();field(p,' TOC \\t "ReportFigure,1" \\h \\z ')
    doc.add_page_break();para('فهرست جدول‌ها',size=18,bold=True,align=WD_ALIGN_PARAGRAPH.RIGHT)
    p=para();field(p,' TOC \\t "ReportTable,1" \\h \\z ')

def body():
    in_refs=False
    for block in (OUT/'report-text.md').read_text(encoding='utf-8').split('\n\n'):
        block=block.strip()
        if not block:continue
        if block.startswith('# '):
            in_refs=block[2:]=='منابع و مراجع'
            chapter(block[2:],'references' if in_refs else 'main')
        elif block.startswith('## '):heading(block[3:])
        elif block.startswith('@table '):main_tables(block.split()[1])
        elif block.startswith('@figure '):
            _,name,title=block.split(' ',2)
            p=para('',align=WD_ALIGN_PARAGRAPH.CENTER);p.paragraph_format.keep_with_next=True
            p.add_run().add_picture(str(OUT/'figures'/f'{name}.png'),width=Cm(16.3));caption(title,True)
        elif block.startswith('@url '):ltr(block[5:],10)
        else:para(block,align=WD_ALIGN_PARAGRAPH.RIGHT if in_refs else None)

def english():
    s=section(None,False)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(OUT/'university-logo.png'),width=Cm(2.7))
    for txt,size,b in [('University of Tehran',18,True),('College of Farabi',16,True),('Faculty of Engineering',16,True),('Department of Computer Engineering',16,True),('Design and Implementation of an AI Clinical Copilot Based on Integration of Patient Context and Validated Medical Knowledge',20,True),('By',14,False),('Erfan Karbalai',16,True),('Supervisor',14,False),('دکتر کاظم فولادی قلعه',16,True),('A Bachelor\'s Project Report in Computer Engineering',16,False),('Submission month and year: ................',14,False)]:
        p=doc.add_paragraph();set_bidi(p,False);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after=Pt(16);p.paragraph_format.space_before=Pt(20 if size==20 else 0)
        if re.search('[\u0600-\u06ff]',txt):set_bidi(p);mixed(p,txt,size,b)
        else:run_format(p.add_run(txt),size,False,b)
    doc.add_page_break();doc.add_paragraph() # compulsory blank back

if __name__=='__main__':
    # Artifact marker is run exactly once immediately before the first authoring.
    marker=OUT/'artifact-marker.json'
    if not marker.exists():
        src=Path('C:/Users/mandegar/.codex/plugins/cache/openai-primary-runtime/documents/26.1011.11113/skills/documents/container_tools/mark_artifact_operation_started.mjs')
        if not src.exists():src=Path(str(src).replace('26.1011.11113','26.1009.11318'))
        proc=subprocess.run(['C:/Program Files/nodejs/node.exe',str(src),'--operation-kind','create','--expected-output-count','1','--output-format','docx'],capture_output=True,text=True)
        if proc.returncode:raise RuntimeError(proc.stderr)
        marker.write_text(json.dumps({'success':True,'path':str(src)}),encoding='utf-8')
    diagrams()
    doc=Document();setup_styles();preliminaries();body()
    from report_appendices import appendices
    appendices(sys.modules[__name__])
    english()
    doc.core_properties.author='عرفان کربالئی';doc.core_properties.title=TITLE
    doc.core_properties.subject='گزارش پروژه کارشناسی؛ Backend و AI'
    doc.core_properties.language='fa-IR'
    doc.save(DOCX)
    (OUT/'sections.json').write_text(json.dumps(chapter_records,ensure_ascii=False,indent=2),encoding='utf-8')
    (OUT/'build-metadata.json').write_text(json.dumps({'paragraphs':len(doc.paragraphs),'tables':len(doc.tables),'figures':len(figures),'text_words':len((OUT/'report-text.md').read_text(encoding='utf-8').split()),'output':str(DOCX)},ensure_ascii=False,indent=2),encoding='utf-8')
    print('Created',DOCX,'tables',len(doc.tables),'paragraphs',len(doc.paragraphs))
