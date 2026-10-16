"""Targeted edits to a copy of the primary DOCX; never rebuild its content."""
from pathlib import Path
from copy import deepcopy
import sys,re,json,hashlib
from PIL import ImageFont
from docx import Document
from docx.shared import Pt,Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
sys.stdout.reconfigure(encoding='utf-8')
OUT=Path(__file__).resolve().parent
ORIGINAL=OUT/'گزارش_پروژه_کارشناسی_عرفان_کربالئی.docx'
REVISED=OUT/'گزارش_پروژه_کارشناسی_عرفان_کربالئی_بازبینی‌شده.docx'
baseline=json.loads((OUT/'review-baseline.json').read_text(encoding='utf-8'))
assert hashlib.sha256(ORIGINAL.read_bytes()).hexdigest()==baseline[str(ORIGINAL)]
d=Document(ORIGINAL);changes=[]
def clean(t):return t.replace('\u200e','').replace('\u00a0',' ')
def setprop(parent,name,value):
    for x in list(parent.findall(qn('w:'+name))):parent.remove(x)
    x=OxmlElement('w:'+name);x.set(qn('w:val'),str(value));parent.append(x);return x
def format_run(r,rtl,size,bold=False,code=False):
    r.font.size=Pt(size);r.font.bold=bold
    font='B Nazanin' if rtl else 'Consolas' if code else 'Times New Roman'
    r.font.name=font
    pr=r._r.get_or_add_rPr()
    rf=pr.find(qn('w:rFonts'))
    if rf is None:rf=OxmlElement('w:rFonts');pr.append(rf)
    for attr in ['ascii','hAnsi','cs','eastAsia']:rf.set(qn('w:'+attr),font)
    setprop(pr,'rtl',1 if rtl else 0);setprop(pr,'szCs',int(size*2));setprop(pr,'bCs',1 if bold else 0)
    lang=pr.find(qn('w:lang'))
    if lang is None:lang=OxmlElement('w:lang');pr.append(lang)
    lang.set(qn('w:val'),'fa-IR' if rtl else 'en-US');lang.set(qn('w:bidi'),'fa-IR')
# No blanket directional characters: explicit Word run properties suffice.
# Parentheses around a wholly English expression belong to that LTR run.
TOKENS=re.compile(r'python[ ]+-m[ ]+[A-Za-z_][A-Za-z0-9_.]*|(?:\.{1,2}/)?\.[A-Za-z][A-Za-z0-9_.-]*|[A-Za-z_][A-Za-z0-9_.]*\([^()\u0600-\u06ff\n]*\)|\([^()\u0600-\u06ff\n]*[A-Za-z][^()\u0600-\u06ff\n]*\)|https?://[^\s]+|[۰-۹]+-[۰-۹]+\.?|(?:الف|ب|پ|ت|ث)-[۰-۹]+\.?|/?[A-Za-z_][A-Za-z0-9_./:{}\[\]?=+%<>#-]*(?:[ \u00a0][A-Za-z][A-Za-z0-9_./:{}\[\]?=+%<>#-]*)*|\d+(?:[.:/][\d]+)*')
def write(p,text,size=14,bold=False):
    for r in list(p._p.findall(qn('w:r'))):p._p.remove(r)
    setprop(p._p.get_or_add_pPr(),'bidi',1)
    pos=0
    for m in TOKENS.finditer(text):
        if m.start()>pos:format_run(p.add_run(text[pos:m.start()]),True,size,bold)
        val=m.group()
        if len(val)<38 and re.search('[A-Za-z]',val):val=val.replace(' ','\u00a0')
        appendix_prefix=bool(re.fullmatch(r'(?:الف|ب|پ|ت|ث)-[۰-۹]+\.?',val))
        format_run(p.add_run(val),appendix_prefix,size if appendix_prefix or re.fullmatch(r'[۰-۹]+-[۰-۹]+\.?',val) else max(size-1,10),bold)
        pos=m.end()
    if pos<len(text):format_run(p.add_run(text[pos:]),True,size,bold)
def replace(i,text,reason):
    p=d.paragraphs[i];old=clean(p.text)
    assert old.strip(),i
    write(p,text)
    changes.append({'paragraph_index':i,'before':old,'after':text,'reason':reason})

EDITS={
251:'عنوان مصوب پروژه «طراحی و پیاده‌سازی دستیار هوشمند بالینی (Clinical AI Copilot) مبتنی بر تلفیق پرونده بیمار و دانش پزشکی معتبر» است. نسخه‌ی موجود، نمونه‌ای دانشگاهی برای ثبت و سازمان‌دهی پرونده، آماده‌سازی متن اسناد و صوت و تولید پاسخ کمکی با مدل آماده است. بخش نخست این هدف در گردش کار برنامه پیاده شده است؛ اتصال به منابع بیرونی دانش پزشکی، از جمله مقالات و دانش‌نامه‌ی دارویی، هنوز در این گردش کار وجود ندارد.',
255:'نمونه‌ی اولیه یا MVP، مخفف Minimum Viable Product، مجموعه‌ای محدود از گردش کارهای قابل اجرا برای پزشک است. ثبت دارو با اقدام کاربر انجام می‌شود و تحلیل ایمنی پیش‌شرط ذخیره‌ی آن نیست. تبدیل رونوشت صوت به یادداشت بالینی نیز به درخواست صریح کاربر نیاز دارد. به این ترتیب ثبت اطلاعات از دسترس‌پذیری مدل جدا می‌ماند و مسئولیت بازبینی خروجی و تصمیم بالینی با پزشک است.',
259:'مسئولیت اصلی رابط کاربری و تجربه‌ی کاربر در پروپوزال با ایمان کربلائی است: داشبورد، نمای بیمار، بارگذاری، نمایش متن، تایم‌لاین، گفتگو و نمایش هشدار و منبع. معماری کلان، قرارداد API، یکپارچه‌سازی، آزمون سرتاسری، استقرار و نمایش نهایی وظایف مشترک‌اند. در این گزارش، جزئیات رابط تا جایی شرح داده می‌شوند که مسیر داده و اتصال آن به قابلیت‌های سمت سرور روشن شود.',
265:'وضعیت هر قابلیت با شواهد آن خوانده می‌شود: رفتارهایی مانند اعتبارسنجی ورودی و قالب درخواست ارائه‌دهنده در اجرای منتخب آزموده شده‌اند؛ برخی مسیرهای دارای کد و آزمون، از جمله APIهای وابسته به پایگاه داده، در این اجرا تأیید نشده‌اند؛ عامل محلی MCP آزمایشی است؛ و استخراج ساختاریافته‌ی خودکار و بازیابی دانش پزشکی بیرونی هنوز پیاده نشده‌اند. نتایج اسناد تسک‌ها با تاریخ و دامنه‌ی همان اجرا تفسیر می‌شوند. شرح پیکربندی Docker نیز در فصل استقرار از نتیجه‌ی اجرای عمومی جدا شده است.',
271:'نمای کلی بیمار برای مرور اولیه، داروهای فعال، حساسیت‌های شدیدتر و نتایج اخیر را با ترتیب مشخص عرضه می‌کند. تایم‌لاین ترتیب رویدادهای ثبت‌شده و منبع هر رویداد را نشان می‌دهد. پاسخ این دو مسیر با queryهای پایگاه ساخته می‌شود: overview داده‌های منتخب را کنار هم می‌گذارد و timeline از رکوردهای اصلی رویداد می‌سازد. تولید تحلیل روایی با مدل، مسیر جداگانه‌ای دارد که برای اسناد پیاده شده است.',
291:'متکی کردن پاسخ به زمینه یا Grounding در این نسخه با خواندن رکوردهای بیمار و قرار دادن آن‌ها در پیام مدل انجام می‌شود. بازیابی یا Retrieval از مقاله و منبع خارجی به مؤلفه‌های دیگری مانند انتخاب قطعه و رتبه‌بندی نیاز دارد که در این مسیر حاضر نیستند. تولید تقویت‌شده با بازیابی یا RAG، مخفف Retrieval-Augmented Generation، در پروپوزال اختیاری است؛ توسعه‌ی آن باید با توجه به نیاز و ارزش افزوده‌ی قابل اندازه‌گیری انجام شود.',
299:'معماری شامل رابط Next.js، سمت سرور FastAPI، پایگاه PostgreSQL و ذخیره‌ساز MinIO است. سرویس‌های مدل و تبدیل گفتار به متن بیرون از این هسته قرار دارند و Backend تماس با آن‌ها را هماهنگ می‌کند. Nginx در پیکربندی استقرار نقطه‌ی ورودی مشترک است. شکل ۳-۱ مسیرهای ارتباط تعریف‌شده در کد و پیکربندی را نشان می‌دهد.',
303:'کد سمت سرور به مسیرهای API و وابستگی‌های آن، طرح‌واره‌های داده، مدل‌های رابطه‌ای، سرویس‌های تخصصی و تنظیمات تقسیم شده است. این اجزا در پوشه‌های api، schemas، models، services و core قرار دارند. مرز لایه‌ها کامل نیست: queryهای برخی منابع در route اجرا می‌شوند و عملیات تخصصی مانند تجمیع زمینه و پردازش سند در service قرار دارند. این سازمان‌دهی پیگیری گردش کار را آسان می‌کند، اما تغییر مشترک سیاست query هنوز به بازبینی هر دو محل نیاز دارد.',
317:'helper دانلود سند در api.ts فقط نشانی مسیر Backend را برمی‌گرداند. EvidenceList این نشانی را در پیوند عادی مرورگر باز می‌کند؛ چنین پیوندی Bearer ذخیره‌شده در localStorage را اضافه نمی‌کند، در حالی که مسیر دانلود Backend آن را می‌خواهد. بنابراین مسیر API و کلیک رابط از نظر انتقال مجوز یکسان عمل نمی‌کنند. آزمون مرورگر باید این اتصال را بررسی کند؛ راهکارهایی مانند دریافت مجوزدار با fetch یا واسط سروری در بخش کارهای آینده مطرح شده‌اند.',
326:'فایل api/v1/api.py پانزده router را ثبت می‌کند: سلامت، ورود با رمز و کد یک‌بارمصرف، کاتالوگ مدل، بیمار و منابع بالینی، اسناد، صوت، گفتگو، دو مسیر ایمنی و عامل. این تجمیع، سطح API را در یک محل قابل بررسی می‌کند. برای تعیین رفتار جاری، فراخوانی واقعی تابع‌ها مبناست؛ بعضی commentها همچنان به مراحل قدیمی توسعه اشاره دارند.',
329:'PatientCreate از PatientBase استفاده می‌کند. نام و نام خانوادگی یک تا صد نویسه دارند. کد ملی اختیاری است و در صورت ارسال باید ده رقم باشد؛ این شرط فقط قالب ورودی را کنترل می‌کند و شامل الگوریتم رقم کنترل نیست. تاریخ تولد PastDate، جنسیت Gender و شماره‌ی تماس رشته‌ای با طول محدود است؛ نشانی نیز متن است. این قیدها خطاهای قراردادی ورودی را پیش از رسیدن به منطق ثبت مشخص می‌کنند.',
369:'واژه‌ی Summary در رابط به نمای ساختاریافته‌ی بیمار اشاره می‌کند. محتوای آن از رکوردها انتخاب می‌شود و برای ساختش مدل فراخوانی نمی‌شود. در مدل داده، ai_summary تحلیل روایی یک سند و cleaned_note متن مشتق از یک ضبط است. این خروجی‌ها به منبع مشخص خود وابسته‌اند و با خلاصه‌ی مستقل کل پرونده تفاوت دارند.',
379:'Context متن خام سند و رونوشت صوت را نیز حمل می‌کند. کاربر برای تبدیل رونوشت به ClinicalNote متن را بازبینی و ارسال می‌کند، اما مدل می‌تواند پیش از این مرحله transcript را در زمینه دریافت کند. پس تأیید یادداشت و دسترسی مدل به متن خام دو مرحله‌ی جدا هستند. افزودن وضعیت بازبینی به قرارداد زمینه می‌تواند به مصرف‌کننده نشان دهد کدام متن هنوز تأیید نشده است.',
406:'خروجی تحلیل سند، متن آزاد ai_summary است. سرویس از این متن رکورد LabResult نمی‌سازد؛ ثبت مقدار آزمایش به طرح‌واره، واحد، منبع و بازبینی جدا نیاز دارد. پروپوزال استخراج فیلدهای منتخب را هدف گرفته بود، در حالی که مسیر فعلی تا استخراج متن و تحلیل سطح سند پیش رفته است. برای نمونه، اشاره‌ی مدل به یک مقدار در خلاصه، آن مقدار را به داده‌ی ساختاریافته‌ی ثبت‌شده تبدیل نمی‌کند.',
409:'رابط اسناد، بارگذاری و فهرست فایل‌ها و وضعیت استخراج و تحلیل را نمایش می‌دهد. آزمون سرتاسری این گردش کار باید انتقال مجوز در دانلود، هدایت ۳۰۷ و دسترسی مرورگر به hostname ذخیره‌ساز را نیز پوشش دهد. شکاف پیوند مستقیم EvidenceList در بخش ۳-۴ توضیح داده شد؛ تست API تنها بخش سمت سرور همین جریان را می‌سنجد.',
423:'OpenRouterProvider قالب درخواست را از OpenAICompatibleProvider به ارث می‌برد و base URL و X-Title را تعیین می‌کند. شناسه‌ی پیش‌فرض مخزن deepseek/deepseek-r1:free است. این مقدار پیکربندی است و دسترس‌پذیری زنده‌ی آن در اجرای حاضر آزموده نشده است. مستندات رسمی OpenRouter قرارداد عمومی API را شرح می‌دهد [۶]؛ انتخاب مدل قابل استفاده باید با حساب و قابلیت مورد نیاز سازگار باشد.',
428:'مقدار gemini-3.6-flash در پیش‌فرض و کاتالوگ ثابت مخزن آمده است. کاتالوگ Gemini از سرویس کشف مدل دریافت نمی‌شود؛ در نتیجه نام‌های آن ممکن است با مدل‌های قابل استفاده‌ی حساب همگام نباشند. صحت نام و دسترس‌پذیری باید پیش از اجرای واقعی بررسی شود و خطای مدل ناموجود از مسیر مدیریت پاسخ ارائه‌دهنده به کاربر برسد.',
446:'رابط گفتگوی بالینی انتخاب‌گر مدل و دو حالت ایمنی و عامل آزمایشی دارد. انتخاب‌ها در localStorage نگهداری می‌شوند و برای پیام ماندگار، provider و model به API ارسال می‌شوند. متن پاسخ و شواهد در مؤلفه‌های جدا نمایش داده می‌شوند. مالکیت اصلی این رابط مطابق پروپوزال با ایمان است و اتصال آن به قراردادهای سمت سرور در حوزه‌ی یکپارچه‌سازی مشترک قرار دارد.',
467:'برای هر قطعه، UnknownValueError باعث عبور از آن قطعه می‌شود. سایر exceptionها در حلقه‌ی تلاش، تا STT_GOOGLE_RETRIES، با پیش‌فرض سه بار، دوباره امتحان می‌شوند و شکست آخر به فراخواننده می‌رسد. اگر فهرست texts خالی بماند، TranscriptionError فارسی ایجاد می‌شود. موفقیت بعضی قطعه‌ها می‌تواند نتیجه‌ای جزئی تولید کند، بدون آنکه fallback یا پرچم «قطعه‌ی حذف‌شده» فعال شود. شرط فعلی، خالی بودن فهرست را می‌سنجد و رشته‌ی خالیِ افزوده‌شده به آن را جدا کنترل نمی‌کند.',
482:'attempt_cleanup خطاهای اجرای complete و ثبت نتیجه را می‌گیرد، اما ساخت provider پیش از آن بلوک انجام می‌شود و در آن محل فقط AIUnavailableError گرفته شده است. exception دیگری از ساخت provider می‌تواند به _process برسد؛ شاخه‌ی عمومی شکست در آن تابع transcript را خالی و transcription_status را failed می‌کند. برای حفظ قطعی رونویسی موفق، آزمون best-effort باید این شاخه را هم پوشش دهد.',
522:'دامنه‌ی داده‌ی ابزارها یکسان نیست. get_patient_documents فقط متادیتای انتخاب‌شده می‌دهد و متن و storage_key را حذف می‌کند؛ get_patient_context خروجی کامل builder، شامل متن استخراج‌شده، را بازمی‌گرداند. بنابراین محدودیت متن فقط به ابزار متادیتای اسناد مربوط است. test_mcp_tools شکل این خروجی‌ها را تعریف می‌کند؛ اجرای زنده‌ی subprocess در بررسی حاضر انجام نشده است.',
539:'کد یک‌بارمصرف از طریق Uniom در پیام‌رسان Soroush Plus ارسال می‌شود. در نبود تنظیمات، route پاسخ ۵۰۳ می‌دهد و جدیدترین رکورد همان شماره را حذف می‌کند؛ تحت رقابت، این رکورد الزاماً همان درخواست فراخواننده نیست. شکست ارسال UniomError به ۵۰۲ تبدیل می‌شود و شماره‌ی مقصد در log خطا ثبت می‌گردد. این رفتار با توضیح ابتدای سرویس مبنی بر ثبت نشدن شماره در log سازگار نیست.',
558:'در بازبینی ۱۴ اکتبر ۲۰۲۶، همان هشت فایل health، CORS، logging، storage، extraction، STT پایه، STT جدید و OpenRouter provider دوباره اجرا شدند. Python 3.12.2 و pytest 8.4.2 استفاده شد؛ وابستگی pytest از زیرپوشه‌ی ابزار گزارش خوانده شد و requirements برنامه تغییر نکرد. ۶۱ مورد جمع‌آوری شد: ۵۹ موفق، صفر failed، دو error در setup و صفر skipped. این آمار از XML و خلاصه‌ی اجرای تازه خوانده شده است.',
559:'هر دو error در fixture فایل test_storage_service.py و پیش از اجرای بدنه‌ی آزمون رخ دادند: اتصال به localhost:9000 رد شد. fixture فقط StorageError را به skip تبدیل می‌کند، ولی خطای اتصال مشاهده‌شده از این شاخه عبور نکرد و error باقی ماند. هفت آزمون استخراج، یازده STT پایه، بیست‌ویک STT جدید و یازده OpenRouter، همراه health و CORS و logging موفق شدند. مدت خلاصه‌ی pytest، ۳۶٫۶۱ ثانیه بود؛ این مقدار زمان اجرای مجموعه‌ی تست است.',
562:'این مجموعه تماس‌های AI و STT را جایگزین می‌کند و PostgreSQL، Docker و گردش کار سرتاسری را اجرا نکرده است. اتصال واقعی OCR در آزمون استخراج بررسی شد و اتصال ذخیره‌ساز به خطای setup رسید. دامنه‌ی ۵۹ نتیجه‌ی موفق به قراردادهای همین فایل‌ها محدود است. فرمان و آرگومان‌های اجرا در گزارش تغییرات و اعتبارسنجی ثبت شده‌اند؛ جزئیات هر مورد نیز در پیوست ث آمده است.',
598:'رابط بالینی این قابلیت‌ها را در گردش کار پزشک به هم متصل می‌کند. سهم اصلی عرفان در هسته‌ی داده و سرویس‌های هوشمند، سهم اصلی ایمان در رابط و سهم مشترک در معماری و یکپارچه‌سازی طبق پروپوزال تفکیک شده است. قرارداد داده و رفتار قابل مشاهده‌ی کد مبنای شرح این سهم‌هاست.',
599:'بازاجرای منتخب در این بازبینی، ۵۹ مورد موفق و دو خطای setup ذخیره‌ساز داشت. جزئیات اجرا در فصل یازدهم ثبت شده است. برای ارزیابی اثر سامانه بر زمان مرور پرونده، خطای درمان یا پیامد بیمار هنوز داده‌ی مرجع و مطالعه‌ی متناسب با آن هدف لازم است. نتیجه‌ی فعلی، ارزیابی محدود قراردادهای نرم‌افزاری است.',
612:'منابع با شماره‌ی یکتا معرفی می‌شوند. ارجاع‌های شماره‌دار مبانی فنی را پشتیبانی می‌کنند؛ مسیر فایل‌ها و نام تابع‌ها شواهد پیاده‌سازی پروژه‌اند. تاریخ دسترسی منابع وب ۱۴ اکتبر ۲۰۲۶ است و گزارش‌های تسک‌های داخلی با دامنه و تاریخ اجرای خود تفسیر می‌شوند.',
853:'جدول نخست تعداد توابع با پیشوند test_ را از AST فایل‌ها نشان می‌دهد. این تعداد با شمار موارد جمع‌آوری‌شده در pytest، به‌ویژه هنگام parametrization، یکسان نیست. جدول دوم نتیجه‌ی بازاجرای منتخب در ۱۴ اکتبر ۲۰۲۶ را ثبت می‌کند. برای آزمون‌های خارج این مجموعه، فقط کد و انتظار رفتاری بررسی شده است.',
850:'فرمان‌های development شامل شروع PostgreSQL و MinIO از ریشه‌ی مخزن و اجرای migration و Uvicorn از پوشه‌ی backend هستند. فرمان‌های build و run بعدی از ریشه اجرا می‌شوند و به فایل Compose استقرار مربوط‌اند. پیش از اجرا باید متغیرهای محیط و hostnameها تنظیم شوند. این فرمان‌ها روی سرور عمومی در بررسی حاضر اجرا نشده‌اند.',
}
for i,text in EDITS.items():replace(i,text,'ویرایش نثر یا تصحیح شواهد اجرایی')
replace(330,'در درخواست‌های PATCH، فیلدها اختیاری‌اند و model_dump(exclude_unset=True) فقط مقادیر ارسال‌شده را اعمال می‌کند. ارسال null با حذف فیلد از درخواست تفاوت دارد: مقادیر ستون‌های nullable قابل پاک شدن‌اند، اما برای ستون غیرnullable ممکن است خطا در مرحله‌ی پایگاه رخ دهد. این تفاوت باید در آزمون طرح‌واره‌های به‌روزرسانی بررسی شود. شرط حداقل طول یک نویسه نیز به‌تنهایی رشته‌ی حاوی فاصله را رد نمی‌کند؛ این کار به strip و اعتبارسنج مناسب نیاز دارد.','حفظ فراخوانی تابع در یک run لاتین و رفع ترکیب نادرست پسوند فارسی با schema')
replace(518,clean(d.paragraphs[518].text),'حفظ فرمان python -m app.mcp_server در یک run لاتین')
replace(798,clean(d.paragraphs[798].text),'اصلاح جهت نام فایل‌های .env و ../.env')

HEADINGS={
314:'۳-۴. مرز رابط کاربری و سمت سرور',
323:'فصل چهارم پیاده‌سازی سمت سرور و مدل داده',
328:'۴-۲. طرح‌واره‌های داده و اعتبارسنجی (Schema)',
347:'۴-۶. نشست پایگاه داده و مهاجرت‌ها',
367:'۵-۳. تفاوت نمای کلی بیمار با زمینه‌ی بالینی',
398:'۶-۴. استخراج متن PDF و تشخیص نوری نویسه‌ها (OCR)',
403:'۶-۵. تحلیل سند با هوش مصنوعی',
410:'۶-۷. آزمون‌های پردازش سند و مسیر توسعه',
421:'۷-۲. درخواست به ارائه‌دهنده‌های OpenAI-compatible و OpenRouter',
425:'۷-۳. قالب پیام و پاسخ در Gemini',
438:'۷-۷. گفتگوی ماندگار و ترتیب ثبت پیام‌ها',
460:'۸-۳. پردازش دسته‌ای صوت در Speechmatics',
464:'۸-۴. آماده‌سازی فایل و تشخیص گفتار با Google web Speech',
469:'۸-۵. ارسال ورودی صوت به OpenRouter',
489:'۹-۱. تغییر جریان ایمنی نسبت به طرح اولیه',
512:'۹-۶. تجزیه‌ی پاسخ و نمایش منبع',
516:'۹-۷. ابزارهای محلی MCP و عامل آزمایشی',
544:'۱۰-۵. نگهداری فایل، داده و گزارش‌های اجرا',
557:'۱۱-۲. بازاجرای آزمون‌های منتخب',
577:'۱۱-۷. توپولوژی استقرار و ترتیب راه‌اندازی',
586:'۱۱-۹. رویه‌ی راه‌اندازی و وضعیت ارزیابی استقرار',
638:'پیوست الف. مرجع مسیرهای رابط برنامه‌نویسی (API)',
698:'پیوست ب. فرهنگ داده و روابط پایگاه',
797:'پیوست پ. تنظیمات و حدود اجرای محیط',
821:'پیوست ت. قراردادهای نمونه و گزیده‌ی کد',
852:'پیوست ث. مرجع آزمون و سناریوهای تکمیلی',
}
for i,p in enumerate(d.paragraphs):
    if p.style.name.startswith('Heading'):
        old=clean(p.text);new=HEADINGS.get(i,old)
        if p.style.name=='Heading 2':
            new=re.sub(r'^([۰-۹]+-[۰-۹]+)(?![۰-۹.])\s+',r'\1. ',new)
            new=re.sub(r'^(الف|ب|پ|ت|ث)ـ([۰-۹]+)\s+',r'\1-\2. ',new)
        size=20 if p.style.name=='Heading 1' else 18
        write(p,new,size,True);setprop(p._p.get_or_add_pPr(),'jc','start')
        if new!=old:changes.append({'paragraph_index':i,'before':old,'after':new,'reason':'تیتر فارسی و شماره‌گذاری یکدست با حفظ شماره'})
        p.paragraph_format.keep_with_next=True
    elif p.style.name=='ReportFigure':
        # Keep captions that are already well rendered; numbering is logical LTR.
        pass
    elif p.style.name=='Code':
        old=p.text
        p.paragraph_format.keep_together=True
        # Preserve every code character. Split Persian string content into runs
        # with a Persian font; do not reverse any token or insert a bidi marker.
        for r in list(p._p.findall(qn('w:r'))):p._p.remove(r)
        setprop(p._p.get_or_add_pPr(),'bidi',0);setprop(p._p.get_or_add_pPr(),'jc','left')
        for part in re.split(r'([\u0600-\u06ff\u200c]+(?:[ ][\u0600-\u06ff\u200c]+)*)',old):
            if part:format_run(p.add_run(part),bool(re.search('[\u0600-\u06ff]',part)),12 if re.search('[\u0600-\u06ff]',part) else 10,code=True)
        assert p.text==old

# Column allocation changes only display. No route/type/setting is renamed.
for ti,t in enumerate(d.tables):
    headers=[clean(c.text) for c in t.rows[0].cells]
    widths=None
    if headers==['ارائه‌دهنده','قالب تماس','کلید و محدودیت']:widths=[5.3,4.5,6.7]
    elif headers[0]=='فعل و مسیر':widths=[6.7,4.1,5.7]
    elif headers[0]=='متغیر':widths=[5.8,5.6,5.1]
    if widths:
        for col,cm in zip(t.columns,widths):col.width=Cm(cm)
        for ri,row in enumerate(t.rows):
            for ci,c in enumerate(row.cells):
                c.width=Cm(widths[ci])
                if headers[0]=='فعل و مسیر' and ci==0 and ri>0:
                    for p in c.paragraphs:
                        for r in p.runs:r.font.size=Pt(9.5)
                if headers[0]=='فعل و مسیر' and ri>0:
                    # Break only at an identifier's CamelCase/underscore or
                    # a path's slash boundary. Letters, keys and order remain.
                    size=9.5 if ci==0 else 10.5
                    f=ImageFont.truetype('C:/Windows/Fonts/times.ttf',round(size*96/72))
                    limit=widths[ci]*72/2.54-12
                    def width(text):return f.getlength(text)*72/96
                    def packed(parts):
                        result=[];line=''
                        for part in parts:
                            if line and width(line+part)>limit:result.append(line);line=part
                            else:line+=part
                        if line:result.append(line)
                        return '\n'.join(result)
                    for p in c.paragraphs:
                        old=clean(p.text)
                        if re.search('[\u0600-\u06ff]',old):continue
                        new=old
                        if ci==0:
                            lines=old.splitlines()
                            new='\n'.join(packed(re.findall(r'/[^/]*',line)) if line.startswith('/') and width(line)>limit else line for line in lines)
                        else:
                            def break_name(m):
                                name=m.group()
                                if width(name)<=limit:return name
                                parts=re.findall(r'[^_]+_+|[^_]+$',name) if '_' in name else re.findall(r'[A-Z]+(?=[A-Z][a-z]|$)|[A-Z]?[a-z0-9]+',name)
                                return packed(parts) if ''.join(parts)==name else name
                            new=re.sub(r'[A-Za-z][A-Za-z0-9_]{17,}',break_name,old)
                        if new!=old:
                            assert re.sub(r'\s','',new)==re.sub(r'\s','',old)
                            for r in list(p._p.findall(qn('w:r'))):p._p.remove(r)
                            setprop(p._p.get_or_add_pPr(),'bidi',0);setprop(p._p.get_or_add_pPr(),'jc','left')
                            format_run(p.add_run(new),False,size)
        changes.append({'table_index':ti,'reason':'تخصیص عرض برای جلوگیری از شکستن نام کلاس، مدل یا متغیر لاتین','widths_cm':widths})
    if headers[0]=='متغیر':
        for row in t.rows[1:]:
            if clean(row.cells[0].text)=='CORS_ORIGINS':
                p=row.cells[1].paragraphs[0];old=clean(p.text)
                assert old=='http://localhost:3000,http://localhost:3100'
                for r in list(p._p.findall(qn('w:r'))):p._p.remove(r)
                setprop(p._p.get_or_add_pPr(),'bidi',0);setprop(p._p.get_or_add_pPr(),'jc','left')
                format_run(p.add_run(old.replace(',',',\n')),False,10.5)
                changes.append({'table_index':ti,'field':'CORS_ORIGINS','reason':'نمایش دو نشانی CSV در دو سطر با حفظ ویرگول و مقادیر'})
    if headers==['نام فیلد','نوع ORM','قید و رابطه','تفسیر']:
        fields={clean(row.cells[0].text):row for row in t.rows[1:]}
        if 'onset_date' in fields and 'name' in fields:
            p=fields['name'].cells[3].paragraphs[0]
            write(p,'نام مشکل بالینی به‌صورت متن آزاد؛ در این جدول قید یکتایی یا کد استاندارد تشخیص ندارد.',12)
            changes.append({'table_index':ti,'field':'name','reason':'تفسیر اختصاصی Condition.name به‌جای توضیح عمومی کاتالوگ دارو'})
        if 'password_hash' in fields and 'role' in fields:
            write(fields['role'].cells[3].paragraphs[0],'نقش حساب: doctor یا admin؛ مبنای سیاست دسترسی پرونده.',12)
            changes.append({'table_index':ti,'field':'role','reason':'توضیح دقیق نقش حساب'})
        if 'conversation_id' in fields and 'role' in fields:
            write(fields['role'].cells[3].paragraphs[0],'نقش پیام در گفتگو: user یا assistant؛ از نقش حساب کاربری مستقل است.',12)
            changes.append({'table_index':ti,'field':'role','reason':'تفکیک نقش پیام از نقش حساب'})

# Keep the TOC field's end marker. Its empty terminal paragraph needs a tiny
# height so it cannot spill onto a header-only page after the last entry.
p=d.paragraphs[246]
if not p.text and p._p.find(qn('w:pPr')+'/'+qn('w:sectPr')) is None and not p._p.xpath('.//w:br'):
    p.paragraph_format.line_spacing=Pt(1)
    p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
    p.paragraph_format.keep_with_next=False;p.paragraph_format.keep_together=False
    for r in p.runs:r.font.size=Pt(1);setprop(r._r.get_or_add_rPr(),'szCs',2)
    changes.append({'reason':'اصلاح ارتفاع پاراگراف پایان field فهرست برای جلوگیری از صفحه‌ی دارای سرصفحه و بدون بدنه'})

# Update repeated chapter titles in header runs without altering PAGE fields.
old_doc=Document(ORIGINAL)
title_map={clean(old_doc.paragraphs[i].text):new for i,new in HEADINGS.items() if old_doc.paragraphs[i].style.name=='Heading 1'}
for sec in d.sections:
    for h in [sec.header,sec.even_page_header,sec.first_page_header]:
        for t in h.tables:
            for row in t.rows:
                for c in row.cells:
                    for p in c.paragraphs:
                        old_header=clean(p.text)
                        if old_header in title_map:write(p,title_map[old_header],10)

# Title and verified administrative information remain exactly as supplied.
# Keep missing dates and the defense placeholder, and leave Bismillah empty.
assert len(d.tables)==baseline['tables']
assert len(d.inline_shapes)==baseline['figures']
d.core_properties.comments='نسخه‌ی بازبینی‌شده؛ اصلاحات نگارشی، جهت متن و راستی‌آزمایی فنی'
d.save(REVISED)
assert hashlib.sha256(ORIGINAL.read_bytes()).hexdigest()==baseline[str(ORIGINAL)]
(OUT/'review-changes.json').write_text(json.dumps(changes,ensure_ascii=False,indent=2),encoding='utf-8')
records=json.loads((OUT/'sections.json').read_text(encoding='utf-8'))
heading1=[p for p in d.paragraphs if p.style.name=='Heading 1']
assert len(records)==len(heading1)
for item,p in zip(records,heading1):item['title']=clean(p.text)
(OUT/'review-sections.json').write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding='utf-8')
print('Revised copy saved:',REVISED)
print('Changed prose paragraphs:',len(EDITS),'Heading/content records:',len(changes),'Tables:',len(d.tables),'Figures:',len(d.inline_shapes))
