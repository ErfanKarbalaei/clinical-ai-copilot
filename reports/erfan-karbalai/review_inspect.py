"""Read-only report inventory and baseline for a separate revised document."""
from pathlib import Path
import sys, json, hashlib, importlib.util, re, subprocess
sys.stdout.reconfigure(encoding='utf-8')
OUT=Path(__file__).resolve().parent
ROOT=OUT.parents[1]
print('Runtime:',sys.executable)
for m in ['docx','pymupdf','PIL','win32com.client']:
    try: print('Module',m, bool(importlib.util.find_spec(m)))
    except ModuleNotFoundError:print('Module',m,False)
from docx import Document
def clean(t):return t.replace('\u200e','').replace('\u00a0',' ')
orig=OUT/'گزارش_پروژه_کارشناسی_عرفان_کربالئی.docx'
d=Document(orig)
paras=[{'index':i,'style':p.style.name,'text':clean(p.text)} for i,p in enumerate(d.paragraphs)]
tables=[[[clean(c.text) for c in row.cells] for row in t.rows] for t in d.tables]
(OUT/'review-original-content.json').write_text(json.dumps({'paragraphs':paras,'tables':tables},ensure_ascii=False,indent=2),encoding='utf-8')
files=[p for p in OUT.glob('*') if p.is_file() and p.suffix in ['.docx','.pdf'] and p.name.startswith('گزارش_پروژه')]
baseline={str(p):hashlib.sha256(p.read_bytes()).hexdigest() for p in files}
baseline['tracked_changes']={p:hashlib.sha256((ROOT/p).read_bytes()).hexdigest() for p in ['.env.example','backend/app/core/config.py','backend/app/services/stt.py','backend/requirements.txt']}
baseline['figures']=len(d.inline_shapes)
baseline['tables']=len(d.tables)
baseline['body_paragraphs']=len(d.paragraphs)
(OUT/'review-baseline.json').write_text(json.dumps(baseline,ensure_ascii=False,indent=2),encoding='utf-8')
heads=[p for p in paras if p['style'].startswith('Heading')]
print('HEADINGS\n'+ '\n'.join(f"{p['index']}: {p['style']} | {p['text']}" for p in heads))
print('Paragraphs:',len(paras),'Tables:',len(tables),'Figures:',len(d.inline_shapes))
for i,t in enumerate(tables):
    if any('test_' in c for row in t for c in row):print('Test table',i,json.dumps(t,ensure_ascii=False))
if '--slice' in sys.argv:
    lo,hi=map(int,sys.argv[-2:])
    print('\n'.join(f"[{p['index']}] {p['text']}" for p in paras[lo:hi]))
