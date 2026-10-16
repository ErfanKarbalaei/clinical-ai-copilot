from pathlib import Path
from docx import Document
import sys

root = Path(__file__).resolve().parent
docx = next(p for p in root.glob("*.docx") if "test" not in p.name)
doc = Document(str(docx))
start, end = int(sys.argv[1]), int(sys.argv[2])
for i in range(start, min(end + 1, len(doc.paragraphs))):
    p = doc.paragraphs[i]
    text = p.text.strip()
    if text:
        print(f"[{i}] [{p.style.name}] {text}")
