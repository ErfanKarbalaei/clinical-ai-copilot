"""Verify preservation and write the separate Persian revision log."""
from pathlib import Path
import json,hashlib,sys,re,zipfile,datetime
from docx import Document
import pymupdf
OUT=Path(__file__).resolve().parent;ROOT=OUT.parents[1]
sys.stdout.reconfigure(encoding='utf-8')
orig=OUT/'گزارش_پروژه_کارشناسی_عرفان_کربالئی.docx'
rev=OUT/'گزارش_پروژه_کارشناسی_عرفان_کربالئی_بازبینی‌شده.docx'
pdf=rev.with_suffix('.pdf');md=OUT/'گزارش_تغییرات_و_اعتبارسنجی.md'
def read(name):return json.loads((OUT/name).read_text(encoding='utf-8'))
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def clean(t):return t.replace('\u200e','').replace('\u00a0',' ')
def compact(t):return re.sub(r'\s+','',clean(t))
base=read('review-baseline.json');meta=read('review-pagination.json')
tests=read('review-test-results.json');technical=read('review-technical-check.json');finaltech=read('review-technical-final.json')
changes=read('review-changes.json');audit=read('review-ooxml-audit.json')
for path,h in base.items():
    if isinstance(h,str):assert sha(Path(path))==h,path
for path,h in base['tracked_changes'].items():assert sha(ROOT/path)==h,path
assert technical['source_hashes']==finaltech['source_hashes'],'Application source changed'
assert technical['code_hashes']==finaltech['code_hashes'],'Code or JSON content changed'
assert finaltech['unsupported_reported_routes']==[] and finaltech['unreported_routes']==[]
assert finaltech['api_count']==finaltech['reported_api_count']==56
assert all(row['matching_orm_tables'] for row in finaltech['orm_field_tables'])
a=Document(orig);b=Document(rev)
assert len(a.tables)==len(b.tables)==56
assert len(a.inline_shapes)==len(b.inline_shapes)==13
assert [p.style.name for p in a.paragraphs if p.style.name.startswith('Heading')]==[p.style.name for p in b.paragraphs if p.style.name.startswith('Heading')]
with zipfile.ZipFile(orig) as az,zipfile.ZipFile(rev) as bz:
    assert bz.testzip() is None
    media={n:sha_data for n in az.namelist() if n.startswith('word/media/') for sha_data in [hashlib.sha256(az.read(n)).hexdigest()]}
    assert media=={n:hashlib.sha256(bz.read(n)).hexdigest() for n in bz.namelist() if n.startswith('word/media/')}
allowed_desc={(t['table_index'],t['field']) for t in changes if 'field' in t and t['field']!='CORS_ORIGINS'}
cell_changes=[]
for ti,(ta,tb) in enumerate(zip(a.tables,b.tables)):
    assert len(ta.rows)==len(tb.rows)
    for ri,(ra,rb) in enumerate(zip(ta.rows,tb.rows)):
        assert len(ra.cells)==len(rb.cells)
        for ci,(ca,cb) in enumerate(zip(ra.cells,rb.cells)):
            if compact(ca.text)!=compact(cb.text):
                field=clean(rb.cells[0].text)
                assert ci==3 and (ti,field) in allowed_desc,(ti,ri,ci)
                cell_changes.append({'table_index':ti,'field':field,'before':clean(ca.text),'after':clean(cb.text)})
render=pymupdf.open(pdf)
assert len(render)==meta['pdf_pages']==meta['word_pages']
assert meta['counts']['main']>=50
assert not render[2].get_text().strip() and not render[17].get_text().strip()
for s in meta['sections']:assert s['physical_page']%2==1 and s['display_page']%2==1
empty_pages=[i+1 for i,p in enumerate(render) if not p.get_text().strip() and not p.get_images()]
overflow=[]
for i,p in enumerate(render,1):
    for w in p.get_text('words'):
        if w[0]<20 or w[2]>p.rect.width-20 or w[1]<20 or w[3]>p.rect.height-20:overflow.append((i,w[4]))
assert not overflow,overflow[:10]
hashes={'original_docx':sha(orig),'original_pdf':sha(orig.with_suffix('.pdf')),'revised_docx':sha(rev),'revised_pdf':sha(pdf)}
result={'preserved_originals':True,'preserved_code_and_json':True,'preserved_figure_binaries':True,'preserved_table_identifiers_types_values_except_three_explanations':True,'changed_explanation_cells':cell_changes,'pages':len(render),'main_pages':meta['counts']['main'],'appendix_pages':meta['counts']['appendix'],'blank_text_and_image_pages':empty_pages,'hashes':hashes,'word_opened_and_rendered':True,'zip_ok':True,'source_hashes_preserved':True}
(OUT/'review-final-audit.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
prose=[c for c in changes if 'before' in c and c.get('reason','').startswith(('ویرایش نثر','حفظ فراخوانی'))]
heads=[c for c in changes if 'تیتر فارسی' in c.get('reason','')]
head_examples=[(c['before'],c['after']) for c in heads if c.get('paragraph_index') in [328,347,367,421,638]]
head_table='\n'.join('| '+x+' | '+y+' |' for x,y in head_examples)
sections='\n'.join(f"| {s['title']} | {s['physical_page']} | {s['display_page']} |" for s in meta['sections'])
cmd='F:\\clinical-ai-copilot\\.venv\\Scripts\\python.exe -B reports\\erfan-karbalai\\review_tests.py'
arglist='\n'.join(str(x) for x in tests['pytest_arguments'])
stats=tests['stats']
text=f'''# گزارش تغییرات و اعتبارسنجی نسخه‌ی بازبینی‌شده

تاریخ بررسی: ۱۴ اکتبر ۲۰۲۶. این نسخه با ویرایش مستقیم کپی DOCX اصلی تهیه شد؛ گزارش از ابتدا بازنویسی نشد. ساختار دوازده فصل، پنج پیوست، منابع، چکیده، جلدها، ۵۶ جدول و ۱۱ نمودار فنی حفظ شد. هیچ فایل اجرایی برنامه تغییر نکرد و commit یا push انجام نشد.

## فایل‌های مبنا و خروجی

- منبع اصلی قابل ویرایش: `{orig}`
- مرجع دیداری قبلی: `{orig.with_suffix('.pdf')}`؛ ۱۲۲ صفحه.
- خروجی Word اصلاح‌شده: `{rev}`
- پیش‌نمایش PDF اصلاح‌شده: `{pdf}`
- گزارش حاضر: `{md}`
- یادداشت قبلی: `reports/erfan-karbalai/گزارش_تطبیق_و_اعتبارسنجی.md`؛ فقط خوانده شد.
- نسخه‌های قدیمی‌تر در `report_builder/output` نیز یافت شدند؛ گزارش تحویل‌شده در پوشه‌ی عرفان مبنای ویرایش بود و نسخه‌های قبلی بازنویسی نشدند.

ورودی‌های رسمی با نام‌های واقعی زیر موجود بودند: پروپوزال `report-inputs/AI_Clinical_Copilot_Proposal_Task_Split_RTL_WORD_FIXED_V3.docx`، راهنما `report-inputs/BSc-project-report-format.pdf` و قالب جلد `report-inputs/BSc-project-report-cover-fa-1401.pdf`. قالب رسمی DOCX جلد در ورودی موجود نیست؛ بازسازی قابل ویرایش قبلی بر مبنای PDF حفظ شد. نام «عرفان کربالئی» طبق تأیید کاربر و نام همکار «ایمان کربلائی» طبق پروپوزال باقی ماند. عنوان مصوب و استاد راهنما دکتر کاظم فولادی قلعه تغییر نکردند.

## اصلاحات نگارشی و جهت متن

{len(prose)} بند نثر یا شرح فنی به‌صورت هدفمند ویرایش شد. توضیح تقسیم لایه‌ها اکنون محل query در route و service را مشخص می‌کند. تکرار توضیح مالکیت رابط، حدود استخراج و شکاف دانلود کاهش یافت و ارجاع داخلی جای توضیح تکراری نشست. محدودیت بالینی، نبود بازیابی خارجی، تفاوت AI-first با قواعد طرح و دامنه‌ی اجرای تست حفظ شدند. تجربه‌ی شخصی یا نتیجه‌ی جدید بالینی اضافه نشد.

نمونه‌ی اصلاح تیترها؛ شماره و موضوع بخش‌ها حفظ شده‌اند:

| پیش از اصلاح | پس از اصلاح |
|---|---|
{head_table}

شماره‌ی بخش‌ها با الگوی یکنواخت و نقطه‌ی جداکننده نمایش داده می‌شود؛ تیترهای پیوست نیز یکدست شدند. عبارت‌های فارسی‌تر، مانند «نشست پایگاه داده» و «نمای کلی بیمار»، در تیتر استفاده شد؛ نام‌های فنی در متن توضیحی باقی ماندند. سرصفحه‌های مربوط به عنوان‌های اصلاح‌شده نیز هماهنگ شدند.

برای پاراگراف فارسی `w:bidi` و برای run فارسی `w:rtl` و زبان `fa-IR` استفاده شد. لاتین‌ها run مستقل LTR دارند. اصلاح جهت به اتکای راست‌چین کردن تنها انجام نشد. به قسمت‌های اصلاح‌شده کاراکترهای نامرئی جهت به‌صورت سراسری اضافه نشد؛ نشانگرهای قدیمی قسمت‌های دست‌نخورده نیز با یک جایگزینی کور حذف نشدند.

موارد دشوار واقعاً در رندر بررسی شدند: `(Schema)` و `(API)`، `model_dump(exclude_unset=True)`، فرمان `python -m app.mcp_server`، فایل‌های `.env` و `../.env`، نام کلاس‌های ارائه‌دهنده، مسیرهای دارای patient_id، نام‌های بلند STT و متن فارسی در رشته‌های JSON. فراخوانی تابع و فرمان MCP اکنون هرکدام یک run لاتین کامل‌اند؛ پرانتز عبارت صرفاً انگلیسی با همان run قرار می‌گیرد.

عرض ستون جدول ارائه‌دهنده افزایش یافت تا `OpenAICompatibleProvider` نشکند. عرض ستون‌های API و تنظیمات بازتخصیص یافت. برای نام‌های دراز API، شکست سطر در مرز CamelCase یا زیرخط و برای مسیرها در مرز slash قرار گرفت؛ ترتیب نویسه‌ها و نام واقعی عوض نشد. مسیرهای API با قلم لاتین ۹٫۵ پوینت و سایر متن فنی جدول عمدتاً ۱۰٫۵ پوینت خوانا هستند. مقدار CSV مربوط به CORS در دو سطر با همان ویرگول نمایش داده شد. متن کد و JSON بدون تغییر نویسه حفظ شد؛ فقط جهت، قلم رشته‌های فارسی و نگه‌داشتن بلوک کد در یک صفحه اصلاح شد.

صفحه‌ی جداکننده‌ی ۱۸ در نسخه‌ی قبلی سرصفحه داشت. پاراگراف خالی پایان field فهرست بدون حذف نشانگر field با آخرین مدخل ادغام شد؛ در رندر نهایی این صفحه کاملاً سفید است. این اصلاح روی کپی گزارش انجام شد و صفحات سفید الزامی مقدمات حذف نشدند.

## راستی‌آزمایی فنی

مبنای کد HEAD برابر `56ed4513818268c320e706a61b6d37e3f401e6da` و تغییرات کاری موجود است. تخصیص Backend/AI به عرفان، Frontend/UX به ایمان و معماری/یکپارچه‌سازی/استقرار به هر دو از پروپوزال حفظ شد. نام نویسنده‌ی commit دلیل مالکیت انحصاری تمام کار فرض نشده است.

| موضوع | نتیجه‌ی بررسی و محل شواهد |
|---|---|
| ثبت API | **۱۵ router** در `backend/app/api/v1/api.py`؛ عدد نادرست ۱۶ در بند ۴-۱ به ۱۵ اصلاح شد |
| پیوست API | هر **۵۶** زوج فعل و مسیر با `application.openapi()['paths']` تطبیق داشت؛ مسیر ساختگی یا جاافتاده یافت نشد |
| داده و مهاجرت | مجموعه‌ی فیلدهای **۱۵ جدول** با `Base.metadata` برابر است؛ هفت فایل revision موجود است |
| تفسیر فیلد | توضیح Condition.name از متن عمومی کاتالوگ دارو به «نام آزاد مشکل، بدون قید یکتایی» اصلاح شد؛ نقش User و ConversationMessage جدا توضیح داده شد؛ نام، نوع و قید فیلدها تغییر نکرد |
| Context و timeline | هفت گروه با سقف ۲۰۰؛ timeline از رکوردها ساخته می‌شود؛ حلقه‌ی مستقل حساسیت ندارد؛ `services/patient_context.py` |
| سند | بیست MiB، استخراج متن با pypdf/OCR، آستانه‌ی مجموع متن و سقف پنجاه صفحه، تحلیل آزاد و سقف ورودی مدل؛ `routes/documents.py` و سرویس‌های extraction و document_analysis؛ استخراج خودکار فیلد بالینی ادعا نشده |
| مدل و گفتگو | سه adapter متن، retry حمل‌ونقل، انتخاب runtime، ثبت user پیش از AI، بیست پیام تاریخچه و evidence JSONB؛ `services/ai.py`، `ai_catalog.py` و `routes/conversations.py` |
| صوت | Speechmatics، Google web Speech و OpenRouter؛ ffmpeg، chunk و fallback از `services/stt.py` خوانده شد؛ متن اکنون گرفتن exception عمومی در retry و امکان پذیرش رشته‌ی خالی را دقیق‌تر شرح می‌دهد |
| پاک‌سازی صوت | exception ساخت provider خارج از بلوک کامل best-effort می‌تواند به شکست کلی و پاک شدن transcript برسد؛ `voice_note_cleanup.py` و `voice_notes.py`؛ نتیجه‌ی آزمون زنده‌ی این شاخه ادعا نشده |
| ایمنی و شواهد | مسیر AI-first و قواعد محدود مستقل‌اند؛ اعتبار ID با حمایت معنایی فرق دارد؛ `ai_medication_safety.py`، `medication_safety.py` و `evidence.py`؛ بازیابی پزشکی بیرونی در جریان جاری نیست |
| MCP و عامل | شش ابزار فقط‌خواندنی؛ ابزار Context متن سند را حمل می‌کند، ابزار documents متادیتا می‌دهد؛ `mcp_server.py` و `services/agent.py`؛ وضعیت آزمایشی حفظ شد |
| ورود و دسترسی | JWT/bcrypt و ACL پرونده، localStorage/cookie قابل خواندن، OTP از Uniom/Soroush، log مقصد و محدودیت شمارش نرخ؛ کد security، dependencies، otp و auth_otp بررسی شد؛ حذف جدیدترین OTP در رقابت الزاماً حذف همان درخواست نیست |
| رابط | انتقال Bearer در api.ts با پیوند عادی EvidenceList یکسان نیست؛ این شکاف در گزارش باقی ماند و از وجود endpoint نتیجه‌ی موفق کلیک گرفته نشد |
| استقرار | دو Dockerfile و پنج سرویس production، Nginx فقط listen 80، نبود ffmpeg در تصویر و اختلاف سقف صوت بررسی شد؛ build و استقرار عمومی در این بازبینی اجرا نشد |

وجود کد و نتیجه‌ی اجرای محدود از ویژگی پیاده‌شده اما ارزیابی‌نشده، مسیر اختیاری/آزمایشی و هدف آینده جدا بیان شده است. هیچ ادعای تشخیص مستقل، تأیید بالینی، انطباق امنیتی، اتصال عملیاتی EHR/HIS یا دانش‌نامه‌ی جامع پزشکی اضافه نشد. منابع و شماره‌های ارجاع قبلی حفظ شدند؛ مرجع، DOI یا benchmark جدیدی ساخته نشد.

## بازاجرای آزمون

فرمان واقعی از ریشه‌ی مخزن:

```text
{cmd}
```

آغاز ثبت UTC: `{tests['verification_utc']}`؛ پایان: `{tests['completed_utc']}`. نسخه‌ها: Python {tests['python']} و pytest {tests['pytest']}. همان هشت فایل قبلی انتخاب شدند. وابستگی‌های pytest و SpeechRecognition از زیرپوشه‌ی ابزار قبلی گزارش خوانده شدند؛ نصب جدید یا تغییر requirements برای سبز کردن تست انجام نشد.

| جمع‌آوری‌شده | موفق | failed | error در setup | skipped |
|---:|---:|---:|---:|---:|
| {stats['collected']} | {stats['passed']} | {stats['failed']} | {stats['errors']} | {stats['skipped']} |

دو مورد `test_upload_presign_download_delete_roundtrip` و `test_storage_key_is_unique_and_sanitized` در fixture MinIO به رد اتصال **localhost:9000** رسیدند؛ بدنه‌ی آن‌ها اجرا نشد. fixture فقط StorageError را skip می‌کند و خطای مشاهده‌شده در آن شاخه گرفته نشد. بنابراین آمار قبلی مستقلاً تأیید شد، اما عبارت «همه‌ی تست‌ها موفق‌اند» درست نیست. نتیجه‌ی اجرای process برابر exit code 1 بود. مدت خلاصه‌ی pytest **۳۶٫۶۱ ثانیه** و مدت suite در XML `{tests['duration']}` ثانیه است؛ ده warning وجود داشت. زمان قبلی ۴۲٫۳۰ ثانیه به این اجرا نسبت داده نشد.

تماس‌های AI/STT جایگزین‌شده بودند؛ آزمون استخراج از OCR موجود نیز استفاده کرد. PostgreSQL fixture، APIهای وابسته به آن، build Docker، سرویس AI زنده و E2E کامل اجرا نشدند. MinIO برای تغییر نتیجه راه‌اندازی نشد. XML جدید پس از بررسی علت خطا از traceback و header احتمالی حساس پاک شد، ولی وضعیت error حفظ گردید. XML قبلی بازنویسی نشد.

آرگومان‌های مؤثر pytest که wrapper ثبت کرده است؛ `-v` از پیکربندی pytest وارد مجموعه شد:

```text
{arglist}
```

## قالب و صفحه‌بندی نهایی

**{len(render)} صفحه‌ی واقعی** در Microsoft Word 16.0 با `ComputeStatistics(wdStatisticPages)` و در PDF با `len(pymupdf.open(...))` یکسان بود. از این تعداد **{meta['counts']['main']} صفحه متن اصلی**، **{meta['counts']['appendix']} صفحه پیوست فنی** و دو صفحه منابع است. مقدمات، فهرست‌ها، جلدها، محل داوری، صفحه‌ی بسم‌الله و سفیدها در حداقل متن اصلی حساب نشده‌اند. حداقل اولیه‌ی ۵۰ صفحه همچنان با متن اصلی تنها برقرار است.

افزایش دو صفحه نسبت به نسخه‌ی اصلی از جریان جدید جدول‌های API و سفید لازم پیش از بخش فرد ناشی شد؛ اندازه‌ی قلم بدنه، حاشیه و فاصله‌ی خط برای افزایش طول تغییر نکرد. A4، متن فارسی چهارده، تیتر بیست/هجده/شانزده، فاصله‌ی ۱۰۰٪، حاشیه‌ی داخلی ۲٫۵ و خارجی ۲، پایین ۲ و آغاز بدنه ۳٫۵ سانتی‌متر، mirrorMargins و ترتیب دانشگاه حفظ شدند.

شروع همه‌ی فصل‌ها، منابع و پیوست‌ها هم در شماره‌ی فیزیکی و هم شماره‌ی نمایشی فرد است. صفحه‌ی اول فصل بدون سرصفحه و با شماره‌ی پایین وسط باقی ماند. فهرست‌ها field واقعی Word دارند و بعد از تغییر تیتر و صفحه‌بندی به‌روزرسانی شدند. روش تفکیک متن اصلی و پیوست، bookmark آغاز بخش و واژه‌های ناحیه‌ی بدنه با مختصات ۹۵ تا ۷۹۵ پوینت بود. فایل متادیتا حدود بخش‌ها را ثبت کرده است؛ تعداد واژه تخمین تعداد صفحه فرض نشد.

| بخش | صفحه‌ی فیزیکی PDF | شماره‌ی نمایشی Word |
|---|---:|---:|
{sections}

## کنترل دیداری و حفظ محتوا

سند واقعاً در Word باز، فهرست‌ها به‌روزرسانی، repaginate، ذخیره و به PDF صادر شد. تمام صفحات به PNG تبدیل و در تصویرهای دوصفحه‌ای با اندازه‌ی اصلی، بدون کوچک‌سازی صفحات، مرور شدند. پس از اصلاحات پایانی، تصویر صفحات تغییرکرده دوباره بررسی شد و یکسان بودن تصاویر دیگر با SHA-256 کنترل گردید. جلد، صفحه‌ی ۳، چکیده، تمام فهرست‌ها، آغاز فصل‌ها، متن مخلوط، جدول‌ها، نمودارها، کد/JSON، منابع و جلد انگلیسی در این مرور پوشش داشتند. بررسی صرفاً استخراج متن یا thumbnail نبود.

صفحه‌ی ۳ کاملاً خالی و مخصوص درج بسم‌الله باقی ماند؛ صفحه‌ی ۱۸ نیز اکنون سفید است. در صفحات بررسی‌شده بریدگی تصویر، جدول بیرون صفحه، حروف فارسی جداافتاده یا وارونگی فرمان اصلاح‌شده دیده نشد. wrap شناسه‌ی بلند در جدول با حفظ ترتیب و مرزهای خوانا انجام می‌شود. کنترل مختصات متن نیز واژه‌ای بیرون فاصله‌ی ایمنی ۲۰ پوینت از لبه‌ی فیزیکی نیافت. سلامت ZIP تأیید شد. شمار OOXML: {audit['persian_paragraphs']} پاراگراف دارای فارسی، {audit['explicit_rtl_paragraphs']} دارای bidi صریح و {audit['sections']} section. بلوک JSON عمداً پاراگراف LTR و run فارسی دارد.

مقایسه‌ی محتوای منطقی جدول‌ها با حذف صرفاً فاصله/شکست سطر، تمام شناسه‌ها، مسیرها، نوع‌ها و مقادیر را برابر یافت؛ فقط سه سلول توضیح فیلد مشخص‌شده اصلاح شدند. تمام بلوک‌های Code و JSON از نظر SHA-256 متن با اصل یکسان‌اند. فایل‌های تصویری داخل DOCX هم بایت‌به‌بایت یکسان ماندند. DOCX و PDF اصلی با hash پیش از بازبینی برابرند.

## تکمیل‌های لازم پیش از تحویل

۱. ماه و سال واقعی تحویل در دو جلد هنوز باید وارد شود. سال تحصیلی ۱۴۰۴–۱۴۰۵ پروپوزال، تاریخ تحویل فرض نشده است.
۲. تصویر بسم‌الله باید در صفحه‌ی ۳ درج شود؛ هیچ توضیح یا caption به آن صفحه افزوده نشد.
۳. فرم رسمی صورت‌جلسه‌ی داوری باید جای صفحه‌ی ۵ را بگیرد. هیچ داور، امضا، تاریخ دفاع، نمره یا تأیید استاد ساخته نشد.
۴. شماره‌ی دانشجویی در منابع نبود؛ اگر دانشگاه آن را می‌خواهد باید توسط دانشجو تکمیل شود. نگارش انگلیسی نام استاد نیز حدس زده نشد و نام فارسی تأییدشده در جلد انگلیسی حفظ شد.
۵. پس از درج فرم/تصویر/تاریخ یا تغییر قلم و نسخه‌ی Word، فهرست و صفحه‌بندی دوباره بررسی شود. PDF حاضر مرجع صفحه‌شماری این نسخه است.
۶. نتیجه‌ی تست کامل پایگاه و MinIO، سرویس زنده و ارزیابی پزشکی هنوز فراهم نشده است. شکاف‌های اجرایی توضیح‌داده‌شده در گزارش با ویرایش سند حل نشده‌اند.

## ابزارها و بهداشت مخزن

برای ویرایش DOCX از Python و python-docx بسته‌ی workspace استفاده شد. رندر با Microsoft Word نصب‌شده و adapter موجود COM/PyMuPDF در Python پروژه انجام شد؛ این کتابخانه‌های رندر در runtime بسته موجود نبودند و نرم‌افزار تازه نصب نشد. marker ویرایش مهارت اسناد با Node بسته اجرا شد. اصل گزارش، ورودی‌ها، یادداشت اعتبارسنجی قبلی و کد برنامه بازنویسی نشدند. helperها و خروجی‌های جدید فقط داخل پوشه‌ی گزارش عرفان‌اند.

خلاصه‌ی ابزارهای بازتولید: `revise_report.py` ویرایش کپی؛ `review_tests.py` بازاجرای محدود؛ `review_technical.py` تطبیق API/ORM؛ `review_render.py` Word/PDF؛ `review_qa_boards.py` تصویرهای بررسی؛ `review_validate.py` کنترل حفظ محتوا و این گزارش. تغییرات ریز در `review-changes.json` و نتیجه‌ی کنترل در `review-final-audit.json` ثبت شده‌اند.

SHA-256 اصل DOCX: `{hashes['original_docx']}`

SHA-256 اصل PDF: `{hashes['original_pdf']}`

SHA-256 DOCX بازبینی‌شده: `{hashes['revised_docx']}`

SHA-256 PDF بازبینی‌شده: `{hashes['revised_pdf']}`
'''
md.write_text(text,encoding='utf-8')
print('Validation saved:',md)
print('Originals preserved. Tables, figures, all code and JSON preserved. Word/PDF',len(render),'main',meta['counts']['main'])
