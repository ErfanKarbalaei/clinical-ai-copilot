from pathlib import Path
import sys
sys.path.insert(0,str(Path(__file__).resolve().parent))
import build_report as K
from docx import Document
import win32com.client,pymupdf
out=Path(__file__).resolve().parent
d=Document()
for align in ['start']:
    p=d.add_paragraph();K.set_bidi(p)
    K.xml(p._p.get_or_add_pPr(),'jc',val=align)
    K.mixed(p,'نمونهٔ فارسی '+align,18)
for token in ['۴ـ۵','۴-۵','4-5']:
    p=d.add_paragraph();K.set_bidi(p);K.xml(p._p.get_or_add_pPr(),'jc',val='start')
    K.run_format(p.add_run('عنوان بخش '),18,True)
    K.run_format(p.add_run('\u200e'+token+'\u200e'),18,False)
d.save(out/'direction-probe.docx')
w=win32com.client.DispatchEx('Word.Application');w.Visible=False;w.DisplayAlerts=0
try:
    x=w.Documents.Open(str(out/'direction-probe.docx'))
    x.ExportAsFixedFormat(str(out/'direction-probe.pdf'),17)
    x.Close(False)
finally:w.Quit()
p=pymupdf.open(out/'direction-probe.pdf');p[0].get_pixmap(matrix=pymupdf.Matrix(1.5,1.5)).save(str(out/'direction-probe.png'))
